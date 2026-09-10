"""Выгрузка аналитики: книга Excel и набор CSV.

Отчёт можно было только смотреть. Чтобы отдать месячные числа руководителю
или в бухгалтерию, их переписывали руками — и переписывали с округлённых
значений на экране, а не с тех, что посчитал сервер.

Здесь тот же отчёт, что и на экране, ложится в таблицы: по листу на разрез
в книге Excel или по файлу в архиве CSV. Разрезы и порядок колонок описаны
одним перечнем `РАЗРЕЗЫ`: и книга, и архив собираются из него, поэтому
разойтись они не могут, а новый разрез добавляется одной строкой.
"""
from __future__ import annotations

import csv
import io
import zipfile
from collections.abc import Callable
from datetime import datetime
from typing import Any

from .errors import ASRHubError
from .logging_setup import get_logger

log = get_logger("analytics.export")

#: Разрез отчёта: ключ в отчёте, имя листа, заголовки и как достать строки.
#:
#: Имя листа не длиннее 31 знака и без символов `[]:*?/\` — это ограничение
#: самого формата, а не вкус: книга с листом «Модели: качество/скорость»
#: не открывается вовсе.
Разрез = tuple[str, str, list[str], Callable[[Any], list[list[Any]]]]


def _пояс() -> str:
    """Смещение часового пояса сервера — для подписи столбца времени."""
    смещение = datetime.now().astimezone().strftime("%z")
    return f"UTC{смещение[:3]}:{смещение[3:]}" if смещение else "UTC"


def МОМЕНТ() -> str:
    """Подпись столбца времени вместе с поясом.

    Excel про часовые пояса не знает вовсе и время со смещением не
    принимает, поэтому в клетку идёт местное время сервера. Молчать об
    этом нельзя: отчёт открывают в другом городе, и «14:35» без пояса —
    это число, к которому нельзя применить ничего.
    """
    return f"Момент ({_пояс()})"


def _ряды(данные: Any, ось: str, ряды: list[str]) -> list[list[Any]]:
    """Ряды по времени — в строки «момент, значение, значение…».

    Момент кладётся датой, а не числом секунд: 1786305554 в таблице — это
    столбец, с которым человек ничего не сделает, а дату Excel сразу
    показывает и умеет по ней группировать.
    """
    метки = (данные or {}).get(ось) or []
    out: list[list[Any]] = []
    for i, метка in enumerate(метки):
        try:
            # Доли секунды у границы корзины — это шум от деления окна на
            # части, а не точность: в клетке они читаются как мусор.
            когда: Any = datetime.fromtimestamp(float(метка)).replace(microsecond=0)
        except (TypeError, ValueError, OSError, OverflowError):
            когда = метка
        строка: list[Any] = [когда]
        for ключ in ряды:
            значения = (данные or {}).get(ключ) or []
            строка.append(значения[i] if i < len(значения) else None)
        out.append(строка)
    return out


#: Столбец выгрузки: подпись по-русски и поле в отчёте.
#:
#: Подписи именно русские: файл уходит бухгалтеру и руководителю, а не
#: разработчику, и «audio_hours» в отчёте для них — это столбец, который
#: сначала надо расшифровать.
Столбец = tuple[str, str]


def _плоско(данные: Any, путь: str) -> Any:
    """Достаёт значение по пути вида «volume.audio_hours»."""
    узел: Any = данные
    for часть in путь.split("."):
        if not isinstance(узел, dict):
            return None
        узел = узел.get(часть)
    return узел


def _таблица(items: Any, столбцы: list[Столбец]) -> list[list[Any]]:
    """Список записей — в строки по перечню столбцов.

    Не тот тип на входе — это ошибка, а не пустая таблица: строка,
    пришедшая вместо списка, разложилась бы по буквам и дала лист из
    пустых клеток. Лист без объяснения хуже отсутствующего.
    """
    записи = items or []
    if not isinstance(записи, (list, tuple)):
        raise TypeError(f"ожидался список записей, пришло {type(записи).__name__}")
    for r in записи:
        if not isinstance(r, dict):
            raise TypeError(f"запись не словарь, а {type(r).__name__}")
    return [[_плоско(r, поле) for _, поле in столбцы] for r in записи]


def _сводка(данные: Any, строки: list[Столбец]) -> list[list[Any]]:
    """Сводка — не таблица записей, а перечень «показатель: значение»."""
    out = []
    for подпись, путь in строки:
        значение = _плоско(данные, путь)
        if значение is not None:
            out.append([подпись, значение])
    return out


_ТОЧНОСТЬ: list[Столбец] = [
    ("", "key"), ("", "jobs"), ("", "words"), ("", "enough"), ("", "wer"), ("", "mer"),
    ("", "wil"), ("", "gap"), ("", "substitutions"), ("", "deletions"),
    ("", "insertions"), ("", "insertion_share"), ("", "wer_avg"),
]

_ЛАТЕНТНОСТЬ: list[Столбец] = [
    ("", "key"), ("", "jobs"), ("", "audio_hours"), ("", "processing_p50"),
    ("", "processing_p95"), ("", "processing_p99"), ("", "rtf_p50"), ("", "rtf_p95"),
    ("", "rtf_p99"), ("", "queue_p50"), ("", "queue_p95"),
]

_ГРУППА: list[Столбец] = [
    ("Значение", "key"), ("Заданий", "jobs"), ("Готово", "completed"),
    ("Ошибок", "failed"), ("Часов звука", "audio_hours"), ("Слов", "words"),
    ("RTF средний", "rtf_avg"),
]

РАЗРЕЗЫ: list[Разрез] = [
    ("overview", "Сводка", ["Показатель", "Значение"],
     lambda d: _сводка(d, [
         ("Период", "period"),
         ("Заданий всего", "jobs.total"),
         ("Готово", "jobs.completed"),
         ("Ошибок", "jobs.failed"),
         ("Отменено", "jobs.cancelled"),
         ("В работе", "jobs.in_progress"),
         ("Из кеша", "jobs.cached"),
         ("Доля успеха", "jobs.success_rate"),
         ("Часов звука", "volume.audio_hours"),
         ("Слов распознано", "volume.words"),
         ("Знаков", "volume.characters"),
         ("Сегментов", "volume.segments"),
         ("RTF средний", "performance.rtf.avg"),
         ("RTF p50", "performance.rtf.p50"),
         ("RTF p90", "performance.rtf.p90"),
         ("Ускорение к реальному времени", "performance.speedup"),
         ("Обработка, средняя, с", "performance.processing_time_s.avg"),
         ("Ожидание в очереди, p95, с", "performance.queue_time_s.p95"),
         ("Средняя уверенность", "quality.confidence.avg"),
         ("WER средний", "quality.wer.avg"),
     ])),
    ("models", "Модели", [
        "Модель", "Название", "Движок", "Заданий", "Готово", "Ошибок",
        "Доля успеха", "Часов звука", "Часов обработки", "RTF средний",
        "RTF p90", "Уверенность", "WER", "Слов"],
     lambda d: _таблица(d, [
         ("", "model"), ("", "name"), ("", "engine"), ("", "jobs"),
         ("", "completed"), ("", "failed"), ("", "success_rate"),
         ("", "audio_hours"), ("", "processing_hours"), ("", "rtf_avg"),
         ("", "rtf_p90"), ("", "confidence_avg"), ("", "wer_avg"), ("", "words")])),
    ("languages", "Языки", [п for п, _ in _ГРУППА], lambda d: _таблица(d, _ГРУППА)),
    ("owners", "Владельцы", [п for п, _ in _ГРУППА], lambda d: _таблица(d, _ГРУППА)),
    ("engines", "Движки", [п for п, _ in _ГРУППА], lambda d: _таблица(d, _ГРУППА)),
    ("tags", "Метки", ["Метка", "Заданий", "Готово", "Ошибок", "Часов звука",
                       "Секунд обработки", "Слов", "RTF"],
     lambda d: _таблица(d, [("", "tag"), ("", "jobs"), ("", "completed"),
                            ("", "failed"), ("", "audio_hours"),
                            ("", "processing_s"), ("", "words"), ("", "rtf")])),
    ("slowest", "Самые медленные",
     ["Номер", "Файл", "Модель", "RTF", "Длительность, с", "Обработка, с"],
     lambda d: _таблица(d, [("", "id"), ("", "filename"), ("", "model"),
                            ("", "rtf"), ("", "duration_s"),
                            ("", "processing_time_s")])),
    ("resources", "Память по моделям",
     ["Модель", "Замеров", "Пик, МБ", "p95, МБ", "Среднее, МБ"],
     lambda d: _таблица((d or {}).get("models"),
                        [("", "model"), ("", "jobs"), ("", "peak_mb"),
                         ("", "p95_mb"), ("", "avg_mb")])),
    ("resources", "Память по загрузке",
     ["Заданий разом", "Замеров", "Пик, МБ", "p95, МБ", "Среднее, МБ"],
     lambda d: _таблица((d or {}).get("concurrency"),
                        [("", "jobs_at_once"), ("", "measurements"),
                         ("", "peak_mb"), ("", "p95_mb"), ("", "avg_mb")])),
    ("resources", "Устройства", ["Устройство", "Заданий", "Часов звука", "RTF"],
     lambda d: _таблица((d or {}).get("devices"),
                        [("", "device"), ("", "jobs"), ("", "audio_hours"),
                         ("", "rtf")])),
    ("audio", "Форматы записей",
     ["Формат", "Заданий", "Часов звука", "Средний размер, МБ"],
     lambda d: _таблица((d or {}).get("formats"),
                        [("", "format"), ("", "jobs"), ("", "audio_hours"),
                         ("", "avg_mb")])),
    ("audio", "Звук на входе",
     ["Речь к шуму", "Записей", "Часов звука", "Уверенность", "Доля низкой уверенности",
      "WER", "Записей с эталоном", "Доля подозрительных сегментов"],
     lambda d: _таблица((d or {}).get("snr_bands"),
                        [("", "key"), ("", "jobs"), ("", "audio_hours"),
                         ("", "confidence_avg"), ("", "low_confidence_share"),
                         ("", "wer_avg"), ("", "wer_jobs"), ("", "suspect_share_avg")])),
    ("queue", "Ожидание в очереди",
     ["Разрез", "Заданий", "p50, с", "p90, с", "p95, с", "p99, с", "Максимум, с"],
     lambda d: _таблица((d or {}).get("by_priority"),
                        [("", "name"), ("", "count"), ("", "p50"), ("", "p90"),
                         ("", "p95"), ("", "p99"), ("", "max")])),
    ("errors", "Ошибки", ["Код", "Заданий", "Сообщение", "Подсказка"],
     lambda d: _таблица((d or {}).get("by_code"),
                        [("", "code"), ("", "count"), ("", "message"),
                         ("", "hint")])),
    ("cache", "Повторы файлов", ["Файл", "Повторов", "Часов звука"],
     lambda d: _таблица((d or {}).get("repeats"),
                        [("", "filename"), ("", "hits"), ("", "audio_hours")])),
    ("timeseries", "Поток заданий",
     [МОМЕНТ(), "Готово", "Ошибки", "Минут звука", "RTF", "Ожидание, с"],
     lambda d: _ряды(d, "labels", ["completed", "failed", "audio_minutes",
                                   "rtf", "queue_time"])),
    ("quality_trend", "Ход качества",
     [МОМЕНТ(), "Уверенность", "WER", "Доля низкой уверенности", "Заданий"],
     lambda d: _ряды(d, "buckets", ["confidence", "wer",
                                    "low_confidence_share", "jobs"])),
    ("accuracy", "Точность по моделям",
     ["Модель", "Записей", "Слов эталона", "Достаточно слов", "WER", "MER", "WIL",
      "WER − MER", "Замен", "Пропусков", "Вставок", "Доля вставок", "WER средний по записям"],
     lambda d: _таблица((d or {}).get("by_model"), _ТОЧНОСТЬ)),
    ("accuracy", "Точность по длительности",
     ["Длительность", "Записей", "Слов эталона", "Достаточно слов", "WER", "MER", "WIL",
      "WER − MER", "Замен", "Пропусков", "Вставок", "Доля вставок", "WER средний по записям"],
     lambda d: _таблица((d or {}).get("by_duration"), _ТОЧНОСТЬ)),
    ("calibration", "Калибровка уверенности",
     ["Корзина от", "до", "Слов", "Средняя уверенность", "Доля верных", "Переоценка"],
     lambda d: _таблица((d or {}).get("bins"),
                        [("", "from"), ("", "to"), ("", "words"), ("", "confidence"),
                         ("", "accuracy"), ("", "gap")])),
    ("latency", "Латентность по моделям",
     ["Модель", "Заданий", "Часов звука", "Обработка p50, с", "p95, с", "p99, с",
      "RTF p50", "RTF p95", "RTF p99", "Очередь p50, с", "Очередь p95, с"],
     lambda d: _таблица((d or {}).get("by_model"), _ЛАТЕНТНОСТЬ)),
    ("latency", "Латентность по длительности",
     ["Длительность", "Заданий", "Часов звука", "Обработка p50, с", "p95, с", "p99, с",
      "RTF p50", "RTF p95", "RTF p99", "Очередь p50, с", "Очередь p95, с"],
     lambda d: _таблица((d or {}).get("by_duration"), _ЛАТЕНТНОСТЬ)),
]


def _таблицы(report: dict[str, Any]) -> list[tuple[str, list[str], list[list[Any]]]]:
    """Отчёт — в набор именованных таблиц. Пустые разрезы пропускаются."""
    out = []
    for ключ, имя, заголовки, достать in РАЗРЕЗЫ:
        try:
            строки = достать(report.get(ключ))
        except Exception as exc:                             # noqa: BLE001
            # Один сломавшийся разрез не должен уносить всю выгрузку: у
            # человека на руках останется файл без одного листа, а не
            # пятисотая ошибка вместо отчёта.
            log.warning("Разрез «%s» не попал в выгрузку: %s", имя, exc)
            continue
        if строки:
            out.append((имя, заголовки, строки))
    return out


#: Готовая таблица выгрузки: имя, заголовки, строки.
Таблица = tuple[str, list[str], list[list[Any]]]


def _собрать_csv(таблицы: list[Таблица], period: str) -> bytes:
    """Архив с файлом CSV на каждую таблицу.

    Разделитель — точка с запятой, кодировка — UTF-8 с меткой порядка байтов.
    И то, и другое ради Excel с русскими настройками: он открывает такой
    файл двойным щелчком и раскладывает по столбцам, а «правильный» CSV с
    запятыми показывает одной колонкой кракозябр.
    """
    буфер = io.BytesIO()
    with zipfile.ZipFile(буфер, "w", zipfile.ZIP_DEFLATED) as архив:
        for имя, заголовки, строки in таблицы:
            текст = io.StringIO()
            писарь = csv.writer(текст, delimiter=";", lineterminator="\r\n")
            писарь.writerow(заголовки)
            писарь.writerows(строки)
            архив.writestr(f"{имя}.csv", "﻿" + текст.getvalue())
        архив.writestr("period.txt", f"Период отчёта: {period}\n")
    return буфер.getvalue()


def to_csv_zip(report: dict[str, Any], period: str) -> bytes:
    """Архив с файлом CSV на каждый разрез аналитики сервера."""
    return _собрать_csv(_таблицы(report), period)


def to_xlsx(report: dict[str, Any], period: str) -> bytes:
    """Книга Excel: по листу на разрез аналитики сервера."""
    return _собрать_xlsx(_таблицы(report),
                         f"ASR Hub — аналитика за период «{period}»")


def _собрать_xlsx(таблицы: list[Таблица], title: str) -> bytes:
    """Книга Excel: по листу на таблицу.

    Отдельно от того, ЧТО выгружается: механику книги — ширину столбцов,
    закреплённую шапку, свойства файла — делят между собой обе выгрузки
    сервера. Две копии этого кода разошлись бы на первой же правке, и
    объяснить человеку, почему у одного отчёта столбцы по содержимому, а у
    второго «####», было бы нечем.
    """
    try:
        from openpyxl import Workbook  # noqa: PLC0415
        from openpyxl.styles import Alignment, Font  # noqa: PLC0415
        from openpyxl.utils import get_column_letter  # noqa: PLC0415
    except ImportError as exc:
        # DependencyMissing описывает отсутствующий движок и подсказывает
        # ставить его скриптом; здесь речь про обычный пакет, и совет должен
        # быть другим. Поэтому ошибка собирается напрямую.
        отказ = ASRHubError(
            "Для выгрузки в Excel нужен пакет openpyxl.",
            hint="Установите его: pip install openpyxl — или выгрузите в CSV, "
                 "он собирается без единого стороннего пакета.",
            cause=exc)
        отказ.code = "dependency_missing"
        отказ.http_status = 503
        raise отказ from exc

    книга = Workbook()
    книга.remove(книга.active)
    шапка = Font(bold=True)
    for имя, заголовки, строки in таблицы:
        лист = книга.create_sheet(имя[:31])
        лист.append(заголовки)
        for клетка in лист[1]:
            клетка.font = шапка
            клетка.alignment = Alignment(vertical="center")
        for строка in строки:
            лист.append(["" if v is None else v for v in строка])
        # Ширина по содержимому: лист со столбцами по умолчанию показывает
        # «####» вместо чисел и обрезанные имена файлов, и первое, что с ним
        # делает человек, — растягивает каждый столбец руками.
        for номер, заголовок in enumerate(заголовки, start=1):
            длины = [len(str(заголовок))]
            длины += [len(str(с[номер - 1])) for с in строки[:200]
                      if с[номер - 1] is not None]
            лист.column_dimensions[get_column_letter(номер)].width = \
                min(52, max(9, max(длины) + 2))
        лист.freeze_panes = "A2"
    if not книга.sheetnames:
        книга.create_sheet("Пусто").append(["За период данных нет"])
    свойства = книга.properties
    свойства.title = title
    свойства.creator = "ASR Hub"
    буфер = io.BytesIO()
    книга.save(буфер)
    return буфер.getvalue()
