"""Регрессии пятого захода ревизии.

Каждая проверка соответствует дефекту, который был воспроизведён на живом
коде: утечка чужих данных, порча результата или зависание. Названия
описывают исходный дефект, а не механику проверки, — чтобы при падении
сразу было понятно, что именно вернулось.
"""
from __future__ import annotations

import builtins
import io
import math
import re
import struct
import subprocess
import sys
import time
import wave
import zipfile
from pathlib import Path
from typing import Any

import pytest
from asrhub.pipeline import audio as audio_mod
from asrhub.pipeline import export as export_mod
from asrhub.pipeline import metrics as metrics_mod
from asrhub.pipeline import postprocess as pp

# ---------------------------------------------------------------------------
# Разграничение доступа
# ---------------------------------------------------------------------------

@pytest.fixture()
def auth_client(tmp_path: Path, monkeypatch: pytest.MonkeyPatch):
    """Сервер с ВКЛЮЧЁННОЙ проверкой ключей.

    Общая фикстура `client` работает без неё — так проще проверять
    остальное, — но разграничение доступа без ключей не проверить: без
    проверки любой запрос считается administratorским.
    """
    from asrhub.api import create_app
    from asrhub.config import load
    from fastapi.testclient import TestClient

    monkeypatch.setenv("ASRHUB_DATA_DIR", str(tmp_path / "data"))
    monkeypatch.setenv("ASRHUB_AUTH_ENABLED", "true")
    monkeypatch.setenv("ASRHUB_MODEL", "demo-simulator")
    monkeypatch.setenv("ASRHUB_ENGINE", "demo")
    app = create_app(load(), start_queue=False)
    with TestClient(app) as test_client:
        admin = (tmp_path / "data" / "api-key.txt").read_text(encoding="utf-8").strip()
        test_client.headers.update({"X-API-Key": admin})
        yield test_client


@pytest.fixture()
def two_users(auth_client):
    """Администратор и два обычных ключа плюс задание первого из них."""
    client = auth_client
    admin = client.headers.get("X-API-Key")
    state = client.app.state.hub

    def make_key(name: str, role: str) -> str:
        response = client.post("/api/keys", json={"name": name, "role": role})
        assert response.status_code == 200, response.text
        return response.json()["key"]

    alice = make_key("alice", "user")
    bob = make_key("bob", "user")
    readonly = make_key("readonly", "readonly")
    job_id = state.db.create_job({
        "filename": "секрет-алисы.wav", "owner": "alice", "status": "running",
        "model": "demo-simulator", "file_path": "/данные/uploads/секрет.wav",
        "text": "ОЧЕНЬ СЕКРЕТНАЯ РАСШИФРОВКА",
        "webhook_url": "https://внутренний.host/hook?token=ТОКЕН",
    })
    return {"admin": admin, "alice": alice, "bob": bob,
            "readonly": readonly, "job_id": job_id}


def test_queue_does_not_leak_other_owners(auth_client, two_users):
    """`GET /api/queue` отдавал чужие задания целиком.

    Прошлый заход закрыл список заданий, а соседний маршрут остался
    открытым: любой ключ видел имена чужих файлов, пути на диске, готовые
    расшифровки и адреса уведомлений вместе с токенами внутри.
    """
    for role in ("bob", "readonly"):
        body = auth_client.get("/api/queue", headers={"X-API-Key": two_users[role]}).text
        assert "секрет-алисы.wav" not in body, f"{role} видит чужой файл"
        assert "/данные/uploads" not in body, f"{role} видит путь на диске"
        assert "ОЧЕНЬ СЕКРЕТНАЯ" not in body, f"{role} видит чужую расшифровку"
        assert "ТОКЕН" not in body, f"{role} видит чужой токен уведомления"

    own = auth_client.get("/api/queue", headers={"X-API-Key": two_users["alice"]}).text
    assert "секрет-алисы.wav" in own, "владелец перестал видеть своё задание"
    admin = auth_client.get("/api/queue", headers={"X-API-Key": two_users["admin"]}).text
    assert "секрет-алисы.wav" in admin, "администратор перестал видеть всё"


def test_job_without_owner_is_not_public(auth_client, two_users):
    """Пустой владелец означал «проверять нечего».

    Задание, созданное клиентом командной строки или ключом без имени,
    читал и удалял кто угодно.
    """
    state = auth_client.app.state.hub
    job_id = state.db.create_job({"filename": "ничей.wav", "owner": "",
                                  "status": "completed", "text": "содержимое"})
    headers = {"X-API-Key": two_users["bob"]}
    assert auth_client.get(f"/api/jobs/{job_id}", headers=headers).status_code == 403
    assert auth_client.delete(f"/api/jobs/{job_id}", headers=headers).status_code == 403
    # Администратору по-прежнему доступно всё.
    assert auth_client.get(f"/api/jobs/{job_id}",
                      headers={"X-API-Key": two_users["admin"]}).status_code == 200


def test_events_hide_administrative_records(auth_client, two_users):
    """`not job_id` пропускало неадминам ровно административные события.

    Создание и отзыв ключей, изменение настроек, загрузка моделей — всё это
    события без job_id, и readonly читал их вместе с именами ключей и ролями.
    """
    state = auth_client.app.state.hub
    state.db.add_event(None, "key_created", "Создан ключ «бухгалтерия» с ролью admin")
    state.db.add_event(None, "queue_paused", "Очередь приостановлена")

    body = auth_client.get("/api/events?limit=50",
                      headers={"X-API-Key": two_users["readonly"]}).text
    assert "бухгалтерия" not in body, "утечка сведений о ключах"
    assert "Очередь приостановлена" in body, "общие события пропали вместе с закрытыми"
    admin_body = auth_client.get("/api/events?limit=50",
                            headers={"X-API-Key": two_users["admin"]}).text
    assert "бухгалтерия" in admin_body


def test_system_paths_are_admin_only(auth_client, two_users):
    """Раскладка файловой системы уходила любому ключу.

    Сама по себе разведка, но именно она превращает прочие находки из
    теоретических в применимые.
    """
    lean = auth_client.get("/api/system", headers={"X-API-Key": two_users["readonly"]}).json()
    assert "paths" not in lean and "database" not in lean
    assert "hardware" in lean, "полезные сведения пропали вместе с закрытыми"
    full = auth_client.get("/api/system", headers={"X-API-Key": two_users["admin"]}).json()
    assert "paths" in full and "database" in full


def test_download_does_not_walk_working_directory(client, tmp_path, monkeypatch):
    """Пустой result_path превращал поиск файла в обход рабочего каталога.

    `Path("")` — это `Path(".")`, поэтому клиенту уходил первый попавшийся
    файл с нужным расширением из каталога, откуда запущен сервер.
    """
    monkeypatch.chdir(tmp_path)
    (tmp_path / "секреты.json").write_text('{"пароль": "СЕКРЕТ"}', encoding="utf-8")
    state = client.app.state.hub
    job_id = state.db.create_job({"filename": "з.wav", "owner": "anonymous",
                                  "status": "completed", "text": "текст",
                                  "result_path": None})
    response = client.get(f"/api/jobs/{job_id}/download?fmt=json")
    assert "СЕКРЕТ" not in response.text
    assert response.status_code != 200 or "секреты.json" not in str(
        response.headers.get("content-disposition", ""))


def test_key_can_be_revoked_with_what_interface_sends(client):
    """Отзыв ключа из интерфейса не работал никогда.

    Интерфейс слал первые шесть символов превью, сервер требовал двенадцать
    и отвечал 400 на любой ключ.
    """
    client.post("/api/keys", json={"name": "лишний", "role": "user"})
    keys = client.get("/api/keys").json()["items"]
    target = next(k for k in keys if k["name"] == "лишний")
    assert len(target["key_id"]) >= 12, "интерфейсу нечего послать"
    assert client.delete(f"/api/keys/{target['key_id']}").status_code == 200
    remaining = [k["name"] for k in client.get("/api/keys").json()["items"]]
    assert "лишний" not in remaining


def test_oversized_upload_is_refused_before_body_is_read(client, monkeypatch):
    """Предел размера проверялся после того, как тело осело на диске.

    FastAPI разбирал multipart целиком, и запрос на десятки гигабайт
    успевал забить временный каталог, прежде чем получить 413.
    """
    import starlette.formparsers as formparsers

    client.app.state.hub.settings.values["max_upload_mb"] = 1
    seen: dict[str, int] = {"bytes": 0}
    original = formparsers.MultiPartParser.parse

    async def spy(self, *args, **kwargs):
        result = await original(self, *args, **kwargs)
        for _, value in result.multi_items():
            seen["bytes"] = max(seen["bytes"], getattr(value, "size", 0) or 0)
        return result

    monkeypatch.setattr(formparsers.MultiPartParser, "parse", spy)
    payload = b"\0" * (8 * 1024 * 1024)
    response = client.post("/api/jobs",
                           files={"file": ("большой.wav", payload, "audio/wav")})
    assert response.status_code == 413
    assert response.json()["code"] == "file_too_large"
    assert seen["bytes"] == 0, "тело всё-таки было принято"


def test_events_are_addressed_to_owner():
    """Рассылка по WebSocket шла всем подряд.

    Владелец подписчика нигде не запоминался, и readonly читал в ленте
    имена чужих файлов и тексты чужих ошибок, а при подключении получал
    ещё и двадцать последних событий.
    """
    from asrhub.api.app import EventHub

    hub = EventHub()
    about_job = {"type": "job.started", "id": "job_1",
                 "filename": "секрет-алисы.wav", "_owner": "alice"}
    common = {"type": "queue_paused"}

    assert hub._visible(about_job, "alice", False) is True
    assert hub._visible(about_job, "bob", False) is False
    assert hub._visible(about_job, "readonly", False) is False
    assert hub._visible(about_job, "любой", True) is True     # администратор
    assert hub._visible(common, "bob", False) is True

    hub._history = [about_job, common]
    assert [m["type"] for m in hub.history_for("bob", False)] == ["queue_paused"]
    assert all("_owner" not in m for m in hub.history_for("alice", False)), \
        "служебное поле утекло клиенту"


# ---------------------------------------------------------------------------
# Корректность результата
# ---------------------------------------------------------------------------

def _wav_with_lead_silence(path: Path, silence_s: float = 3.0,
                           total_s: float = 5.0, rate: int = 16000) -> Path:
    frames = bytearray()
    for index in range(int(rate * total_s)):
        second = index / rate
        value = 0.0 if second < silence_s else 0.4 * math.sin(2 * math.pi * 320 * second)
        frames += struct.pack("<h", int(value * 32767))
    with wave.open(str(path), "wb") as handle:
        handle.setnchannels(1)
        handle.setsampwidth(2)
        handle.setframerate(rate)
        handle.writeframes(bytes(frames))
    return path


@pytest.mark.skipif(not audio_mod.has_ffmpeg(), reason="нужен ffmpeg")
def test_trimmed_silence_does_not_move_timestamps(tmp_path: Path):
    """Обрезка начальной тишины сдвигала все таймкоды, и это никто не учитывал.

    Настройка включена по умолчанию, поэтому субтитры любой записи с
    паузой в начале уезжали ровно на длину этой паузы.
    """
    source = _wav_with_lead_silence(tmp_path / "тишина.wav", silence_s=3.0)
    settings = {"audio_trim_silence": True, "audio_sample_rate": 16000,
                "audio_channels": "mono"}
    prepared = audio_mod.prepare(source, tmp_path / "work", settings)

    assert prepared.shifted, "сдвиг не замечен"
    assert 2.5 <= prepared.offset_s <= 3.5, f"сдвиг измерен неверно: {prepared.offset_s}"
    # Речь начинается в нуле подготовленного файла и в трёх секундах исходного.
    assert 2.5 <= prepared.to_source_time(0.0) <= 3.5


@pytest.mark.skipif(not audio_mod.has_ffmpeg(), reason="нужен ffmpeg")
def test_no_shift_without_trimming(tmp_path: Path):
    source = _wav_with_lead_silence(tmp_path / "тишина.wav")
    prepared = audio_mod.prepare(source, tmp_path / "work", {
        "audio_trim_silence": False, "audio_sample_rate": 16000,
        "audio_channels": "mono"})
    assert prepared.offset_s == 0.0 and not prepared.shifted


def test_speed_change_is_accounted_for():
    """Изменение темпа сжимало таймкоды, и коэффициент нигде не применялся."""
    prepared = audio_mod.Prepared(channels=[], offset_s=2.0, speed=1.25)
    assert prepared.to_source_time(0.0) == 2.0
    assert prepared.to_source_time(4.0) == 7.0


@pytest.mark.parametrize("source,expected", [
    ("Осталось три, четыре дня.", "Осталось 3, 4 дня."),
    ("Пять, шесть человек ждут.", "5, 6 человек ждут."),
    ("Один, два, три, поехали!", "1, 2, 3, поехали!"),
    ("Ему двадцать.", "Ему 20."),
    ("Код три три три.", "Код три три три."),
    # То, что должно нормализоваться, — по-прежнему нормализуется.
    ("Двадцать пять рублей.", "25 ₽."),
    ("Сто двадцать три.", "123."),
    ("Две тысячи двадцать четыре год", "2024 год"),
    ("Скидка тридцать процентов", "Скидка 30 %"),
])
def test_numbers_do_not_merge_across_punctuation(source, expected):
    """Числительные складывались через запятую, а знак препинания терялся.

    «три, четыре дня» превращалось в «7 дня», «один, два, три, поехали!» —
    в «6 поехали!». Нормализация включена по умолчанию и работает после
    расстановки пунктуации, то есть портила именно то, что модель только
    что расставила.
    """
    assert pp._builtin_itn_ru(source) == expected


@pytest.mark.parametrize("text", [
    "Он любит хлебать горячий суп",
    "Нахлебался чаю",
    "Расхлебали кашу",
    "Сукно на столе",
])
def test_profanity_filter_leaves_ordinary_words(text):
    """Корень искался в любом месте слова, и «хлебать» становилось «х******»."""
    assert pp.filter_profanity(text, "mask") == (text, 0)


@pytest.mark.parametrize("text", ["Иди на хуй", "Полная хуйня", "Заебал уже",
                                  "Это блядство"])
def test_profanity_filter_still_works(text):
    result, hits = pp.filter_profanity(text, "mask")
    assert hits >= 1 and result != text


@pytest.mark.parametrize("seconds,expected", [
    (59.9996, "00:01:00,000"),
    (3599.9996, "01:00:00,000"),
    (59.9994, "00:00:59,999"),
    (0.0, "00:00:00,000"),
])
def test_timestamps_carry_over_correctly(seconds, expected):
    """Округление добавляло секунду, но перенос дальше не шёл.

    Получались «00:00:60,000» и «00:59:60,000» — такой блок строгие плееры
    просто отбрасывают.
    """
    assert export_mod.format_timestamp(seconds, "srt") == expected


def test_subtitle_wrapping_keeps_every_word():
    """Лишние строки отрезались вместе со словами."""
    phrase = ("Мы обсудили условия поставки и договорились перенести отгрузку "
              "на следующий понедельник, потому что склад закрыт")
    wrapped = export_mod.wrap_subtitle(phrase, 42, 2)
    assert wrapped.replace("\n", " ").split() == phrase.split()


def test_ass_escapes_braces():
    """В ASS фигурные скобки — команды оформления.

    «Скидка {30} процентов» отрисовывалась без числа: всё в скобках
    считалось командой и не показывалось.
    """
    result = export_mod.to_ass(
        {"segments": [{"start": 0.0, "end": 2.0, "text": "Скидка {30} процентов"}]}, {})
    assert r"\{30\}" in result
    assert "{30}" not in result.replace(r"\{30\}", "")


def test_speaker_labels_setting_applies_to_text():
    """`include_speaker_labels` не действовала на txt.

    Хуже того, тот же текст шёл в расчёт точности, и каждая реплика
    добавляла к эталону две лишние вставки, завышая WER.
    """
    segments = [{"start": 0.0, "end": 2.0, "text": "Первая", "speaker": "Оператор"},
                {"start": 2.0, "end": 4.0, "text": "Вторая", "speaker": "Клиент"}]
    result = {"segments": segments, "text": "", "speakers": ["Оператор", "Клиент"]}
    without = export_mod.to_txt(result, {"include_speaker_labels": False})
    assert "Оператор:" not in without and "Первая" in without
    with_labels = export_mod.to_txt(result, {"include_speaker_labels": True})
    assert "Оператор:" in with_labels


def test_accuracy_does_not_hang_on_long_texts():
    """Посимвольный расчёт был квадратичным без ограничителя.

    Часовая расшифровка считалась больше получаса, и всё это время задание
    висело на 96 % и не отменялось.
    """
    vocabulary = ["сегодня", "мы", "обсуждали", "условия", "поставки", "сроки", "отгрузки", "договор", "оплата", "склад", "менеджер", "клиент", "заявка", "счёт", "документы"]
    import random
    random.seed(7)
    reference = [random.choice(vocabulary) for _ in range(8000)]   # примерно час
    hypothesis = list(reference)
    for _ in range(800):
        hypothesis[random.randrange(len(hypothesis))] = random.choice(vocabulary)

    started = time.time()
    result = metrics_mod.detailed(" ".join(reference), " ".join(hypothesis))
    elapsed = time.time() - started
    assert elapsed < 30, f"разбор занял {elapsed:.1f} с"
    assert 0.05 < result["wer"] < 0.2, result["wer"]
    assert 0.0 < result["cer"] < 0.2, result["cer"]


def test_chunked_distance_matches_direct_computation():
    """Ускорение не должно менять числа."""
    import random
    random.seed(11)
    vocabulary = ["раз", "два", "три", "четыре", "пять", "шесть", "семь", "восемь", "девять", "десять"]
    reference = [random.choice(vocabulary) + str(i) for i in range(600)]
    hypothesis = list(reference)
    for _ in range(60):
        hypothesis[random.randrange(len(hypothesis))] = "ошибка"
    direct = metrics_mod._levenshtein(reference, hypothesis)
    chunked = metrics_mod._levenshtein_chunked(reference, hypothesis)
    assert direct.error_rate == chunked.error_rate


def test_glossary_error_does_not_lose_transcript():
    """Ошибка в строке замены летела мимо перехвата.

    `subn` разбирает замену как шаблон, и ссылка на несуществующую группу
    обрушивала весь шаг постобработки вместе с расшифровкой.
    """
    text, count = pp.apply_glossary("версия три", {"re:версия": "\\1 версия"})
    assert text == "версия три" and count == 0
    # Исправные правила продолжают работать.
    assert pp.apply_glossary("версия три", {"re:версия": "редакция"}) == ("редакция три", 1)


def test_waveform_reads_both_segment_shapes():
    """Полоса строится после постобработки, где сегменты — словари.

    Обращение через getattr к словарю молча давало None, и кривые по
    говорящим переставали строиться вовсе.
    """
    from asrhub.engines.base import Segment
    from asrhub.pipeline import waveform as waveform_mod

    objects = [Segment(0.0, 2.0, "раз", speaker="Оператор"),
               Segment(2.0, 4.0, "два", speaker="Клиент")]
    assert waveform_mod._field(objects[0], "speaker") == "Оператор"
    assert waveform_mod._field(objects[0].to_dict(), "speaker") == "Оператор"


# ---------------------------------------------------------------------------
# Аналитика, очередь, хранение
# ---------------------------------------------------------------------------

@pytest.fixture()
def analytics_db(tmp_path: Path):
    from asrhub.analytics import Analytics
    from asrhub.db import Database

    database = Database(tmp_path / "a.sqlite3")
    moment = time.time()
    for owner, language, engine, status, count in (
            ("alice", "ru", "whisper", "completed", 2),
            ("bob", "en", "vosk", "completed", 4),
            ("alice", "ru", "whisper", "failed", 1)):
        for _ in range(count):
            job_id = database.create_job({
                "owner": owner, "language": language, "engine": engine, "model": "m",
                "media_duration_s": 60.0, "created_at": moment - 60})
            database.update_job(job_id, status=status, finished_at=moment - 30,
                                processing_time_s=10.0, rtf=0.16)
    return Analytics(database)


def test_analytics_sections_respect_owner(analytics_db):
    """`by_language` и `by_engine` не передавали владельца в `_group`.

    Соседние методы передавали — то есть это была опечатка, из-за которой
    обычный ключ видел сводку по всему серверу.
    """
    languages = analytics_db.by_language("day", owner="alice")
    engines = analytics_db.by_engine("day", owner="alice")
    assert {row["key"] for row in languages} == {"ru"}
    assert {row["key"] for row in engines} == {"whisper"}
    assert sum(row["jobs"] for row in languages) == 3


def test_failure_rate_uses_one_population(analytics_db):
    """Числитель считался по владельцу, знаменатель — по всему серверу."""
    report = analytics_db.errors("day", owner="alice")
    assert report["total_failed"] == 1
    assert report["total_jobs"] == 3
    assert abs(report["failure_rate"] - 1 / 3) < 0.001


def test_queue_is_not_starved_by_pending_retries(tmp_path: Path):
    """Отбор «время повтора наступило» шёл после LIMIT.

    Больше пятисот заданий в статусе retry забивали окно предвыборки
    целиком, и воркеры простаивали, хотя готовые задания были.
    """
    from asrhub.db import Database

    database = Database(tmp_path / "q.sqlite3")
    moment = time.time()
    for index in range(600):
        job_id = database.create_job({"status": "retry", "media_duration_s": 60.0,
                                      "priority": 90})
        database.update_job(job_id, status="retry", queued_at=moment + 300,
                            created_at=moment - 1000 + index)
    for _ in range(5):
        job_id = database.create_job({"status": "queued", "media_duration_s": 60.0,
                                      "priority": 50})
        database.update_job(job_id, status="queued", queued_at=moment - 10)

    for order in ("created_at ASC", "media_duration_s ASC", "priority DESC"):
        ready = database.list_jobs(status=["queued", "retry"], limit=500, order=order,
                                   light=True, ready_before=time.time())
        assert len(ready) == 5, f"политика {order}: очередь встала"


def test_cleanup_is_bounded_and_covers_unfinished(tmp_path: Path):
    """Уборка выбирала всё разом и не трогала незавершённые задания.

    Понижение срока хранения вытаскивало в память сотни тысяч строк и
    занимало блокировку записи на десятки минут, а файлы заданий, застрявших
    в очереди, не удалялись никогда.
    """
    from asrhub.db import CLEANUP_BATCH, Database

    database = Database(tmp_path / "c.sqlite3")
    moment = time.time()
    ancient = moment - 400 * 86400
    for _ in range(10):
        job_id = database.create_job({"status": "completed"})
        database.update_job(job_id, status="completed", finished_at=ancient)
    for _ in range(5):
        job_id = database.create_job({"status": "queued"})
        database.update_job(job_id, created_at=ancient)
    for _ in range(3):
        database.create_job({"status": "queued"})       # свежие — трогать нельзя

    removed = database.cleanup(results_days=30)
    assert removed["jobs"] == 15
    assert database.count_jobs() == 3
    assert CLEANUP_BATCH > 0
    assert "LIMIT" in _cleanup_sql(), "предел на заход исчез из запроса"


def _cleanup_sql() -> str:
    source = Path(__file__).resolve().parent.parent / "server" / "asrhub" / "db.py"
    text = source.read_text(encoding="utf-8")
    start = text.index("def cleanup(")
    return text[start:start + 2000]


# ---------------------------------------------------------------------------
# Сценарии установки
# ---------------------------------------------------------------------------

def test_gpu_driver_state_reads_installed_driver(repo_root: Path):
    """«A || B && C» разбиралось как «(A || B) && C».

    При наличии nvidia-smi без модуля в updates/dkms состояние «драйвер
    стоит, нужна перезагрузка» читалось как «драйвера нет», и скрипт шёл
    ставить драйвер заново.
    """
    script = f"""
        set -o errexit -o nounset -o pipefail
        source "{repo_root}/scripts/lib/common.sh"
        source "{repo_root}/scripts/lib/detect.sh"
        source "{repo_root}/scripts/lib/gpu.sh"
        have() {{ [[ "$1" == nvidia-smi ]]; }}
        nvidia-smi() {{ return 1; }}
        gpu_driver_state 0x10de
    """
    result = subprocess.run(["bash", "-c", script], capture_output=True, text=True,
                            timeout=60, env={"ASRHUB_QUIET": "1", "PATH": "/usr/bin:/bin"})
    assert result.stdout.strip() == "installed-noload", result.stdout


def test_installer_does_not_copy_over_itself(repo_root: Path, tmp_path: Path):
    """Запуск из установленной копии сносил установку.

    Цикл удалял каталог и тут же копировал его сам в себя, оставляя
    установку без server — при том что шапка скрипта обещает обратное.
    """
    import shutil

    copy = tmp_path / "установка"
    shutil.copytree(repo_root, copy, ignore=shutil.ignore_patterns(
        ".git", "build", "__pycache__", ".pytest_cache", ".ruff_cache", "*.whl"))
    result = subprocess.run(
        ["bash", "scripts/install.sh", "--dry-run", "--no-interactive", "--yes",
         "--profile", "light", "--skip-models", "--no-service",
         "--prefix", str(copy), "--data", str(tmp_path / "data")],
        cwd=copy, capture_output=True, text=True, timeout=300)
    assert result.returncode == 0, result.stdout[-2000:]
    assert "копирование пропущено" in result.stdout
    assert (copy / "server").is_dir() and (copy / "scripts").is_dir()


def test_doctor_works_from_any_directory(repo_root: Path, tmp_path: Path):
    """Путь к пакету был относительным, а doctor.sh запускают откуда угодно.

    Установщик прямо предлагает «bash /opt/asrhub/scripts/doctor.sh», и на
    исправной установке весь раздел движков объявлялся сломанным.
    """
    import shutil

    prefix = tmp_path / "установка"
    (prefix / "venv" / "bin").mkdir(parents=True)
    (prefix / "venv" / "bin" / "python").symlink_to(sys.executable)
    shutil.copytree(repo_root / "server", prefix / "server")
    elsewhere = tmp_path / "другой-каталог"
    elsewhere.mkdir()

    result = subprocess.run(
        ["bash", str(repo_root / "scripts" / "doctor.sh"), "--prefix", str(prefix)],
        cwd=elsewhere, capture_output=True, text=True, timeout=180)
    assert "не удалось выполнить" not in result.stdout, result.stdout[-1500:]
    assert "demo" in result.stdout


def test_env_file_format_suits_systemd(repo_root: Path):
    """`export` в env.sh ломал EnvironmentFile.

    systemd читает строго ИМЯ=ЗНАЧЕНИЕ и строку с «export » отбраковывает —
    собранный whisper.cpp сервер попросту не находил.
    """
    text = (repo_root / "scripts" / "lib" / "whispercpp.sh").read_text(encoding="utf-8")
    assert "printf 'export ASRHUB_WHISPER_CPP" not in text
    assert "printf 'ASRHUB_WHISPER_CPP=%s" in text


def test_powershell_confirm_respects_default(repo_root: Path):
    """`Confirm-Action` без консоли отвечала «да» на всё.

    `uninstall.ps1 -Purge` из задачи планировщика удалял каталог данных с
    базой и моделями, ни о чём не спросив, — при умолчании «нет».
    """
    text = (repo_root / "scripts" / "lib" / "Common.psm1").read_text(encoding="utf-8")
    body = text.split("function Confirm-Action")[1].split("\n}")[0]
    assert "return $true" not in body.split("UserInteractive")[1].split("\n")[0], \
        "безусловное согласие вернулось"
    assert "$Default -eq 'y'" in body


def test_powershell_rollback_captures_each_directory(repo_root: Path):
    """Блок отката связывался с переменной поздно.

    К моменту вызова $dir равнялся последнему значению цикла, и откат пять
    раз удалял data\\tmp, оставляя каталог программы с venv на диске.
    """
    text = (repo_root / "scripts" / "install.ps1").read_text(encoding="utf-8")
    assert "GetNewClosure()" in text
    assert "$captured = $dir" in text


def test_docker_uid_is_not_root_under_sudo(repo_root: Path):
    """Под sudo id -u давал ноль, и контейнер работал от root, минуя gosu."""
    text = (repo_root / "scripts" / "install.sh").read_text(encoding="utf-8")
    assert "ASRHUB_UID=${SUDO_UID:-$(id -u)}" in text
    assert "ASRHUB_GID=${SUDO_GID:-$(id -g)}" in text


def test_service_user_is_created(repo_root: Path):
    """Служба systemd по умолчанию работала от root."""
    text = (repo_root / "scripts" / "install.sh").read_text(encoding="utf-8")
    assert "useradd --system --no-create-home" in text
    assert 'SERVICE_USER="asrhub"' in text


def test_windows_service_does_not_leave_broken_registration(repo_root: Path):
    """sc.exe регистрировал python.exe как службу — она не может стартовать.

    Диспетчер возвращал ошибку 1053, а служба оставалась зарегистрированной
    с автозапуском и тремя попытками перезапуска при каждой загрузке.
    """
    text = (repo_root / "scripts" / "service.ps1").read_text(encoding="utf-8")
    assert "sc.exe create" not in text, "сломанная регистрация службы вернулась"
    assert "Install-AsTask" in text


def test_web_interface_regressions(repo_root: Path):
    """Правки в интерфейсе, которые нечем проверить кроме как по коду."""
    app_js = (repo_root / "server" / "asrhub" / "web" / "app.js").read_text(encoding="utf-8")
    # Журнал: 403 на /api/logs гасил и панель событий
    assert "Promise.allSettled" in app_js
    # Событие не перерисовывает раздел целиком, стирая фильтры
    assert "function refreshLiveViews" in app_js
    # Файлы, добавленные во время отправки, больше не пропадают
    assert "const batch = state.files.slice()" in app_js
    # Отзыв ключа шлёт то, что сервер принимает
    assert "k.key_id" in app_js
    # У кнопок-иконок есть доступное имя
    assert 'aria-label="Закрыть"' in app_js

    charts_js = (repo_root / "server" / "asrhub" / "web" / "charts.js").read_text(
        encoding="utf-8")
    assert "createTextNode(item.name)" in charts_js, "имя в легенде снова вставляется как HTML"


def test_light_listing_keeps_fields_the_interface_needs(auth_client):
    """Облегчённый список обязан нести прогресс и стадию.

    `/api/queue` перевели на light вместе с сужением по владельцу, а полосу
    выполнения на главном экране интерфейс берёт именно оттуда: без этих
    двух полей она замерла бы на нуле.
    """
    state = auth_client.app.state.hub
    job_id = state.db.create_job({"filename": "з.wav", "status": "running",
                                  "owner": "ключ"})
    state.db.update_job(job_id, status="running", progress=0.42, stage="распознавание")
    item = auth_client.get("/api/queue").json()["items"][0]
    assert item["progress"] == 0.42 and item["stage"] == "распознавание"
    assert "text" not in item, "облегчённый список снова тянет расшифровку"


# ---------------------------------------------------------------------------
# Загрузка GigaAM: куда пишем, что качаем и как называем причину отказа
# ---------------------------------------------------------------------------


РЕАЛЬНЫЙ_ОТКАЗ = [
    "v3_e2e_rnnt: OSError: [Errno 30] Read-only file system: '/home/asrhub'",
    "ai-sage/GigaAM-v3: ValueError: Model 'ai-sage/GigaAM-v3' not found. "
    "Available model names: ['ctc', 'rnnt', 'e2e_ctc', 'e2e_rnnt', 'ssl', 'emo', "
    "'v1_ctc', 'v1_rnnt', 'v1_ssl', 'v2_ctc', 'v2_rnnt', 'v2_ssl', 'v3_ctc', "
    "'v3_rnnt', 'v3_e2e_ctc', 'v3_e2e_rnnt', 'v3_ssl', 'multilingual_ctc']",
    "transformers: InstantiationException: Error in call to target "
    "'modeling_gigaam.GigaAMASR': RuntimeError('Tensor on device cpu is not on "
    "the expected device meta!')",
]


def test_a_model_name_in_the_error_does_not_become_a_diagnosis():
    """«ssl» в перечне доступных моделей выдавалось за отказ сети.

    Отказ был из-за прав на запись, а разбор нашёл в тексте подстроку «ssl» —
    внутри имени модели «ssl» из перечня, который печатает вторая попытка, —
    и объявил, что сервер не достучался до хранилища весов. Человек ушёл
    проверять интернет, а дело было в каталоге.
    """
    from asrhub.engines.gigaam_engine import _load_failure

    отказ = _load_failure("gigaam-v3-e2e-rnnt", РЕАЛЬНЫЙ_ОТКАЗ, "cuda",
                          "/var/lib/asrhub/models")
    текст = str(отказ)
    assert "хранилищу весов" not in текст, "сеть по-прежнему назначена виноватой"
    assert "только для чтения" in текст, текст


def test_the_failing_path_is_named():
    """«Куда именно не удалось записать» — самое полезное в таком отказе."""
    from asrhub.engines.gigaam_engine import _load_failure

    отказ = _load_failure("gigaam-v3-e2e-rnnt", РЕАЛЬНЫЙ_ОТКАЗ, "cuda", "")
    # Путь есть и в общем списке попыток — но там он тонет. Нужна своя строка.
    строки = [s for s in отказ.hint.splitlines() if s.startswith("Путь,")]
    assert строки, "отдельной строки с путём нет:\n" + отказ.hint
    assert "/home/asrhub" in строки[0], строки


def test_the_first_attempt_decides_the_reason():
    """Первая попытка — штатный путь; остальные падают по своим поводам.

    Разбор шёл по склейке всех попыток, и текст запасных мог перебить
    причину из первой. Порядок попыток не случаен, и разбор обязан его
    уважать.
    """
    from asrhub.engines.gigaam_engine import _load_failure

    отказ = _load_failure("m", [
        "v3_rnnt: OSError: [Errno 28] No space left on device",
        "transformers: ConnectionError: Connection refused",
    ], "cpu", "")
    assert "места на диске" in str(отказ), str(отказ)


def test_a_tokenizer_is_not_mistaken_for_a_token():
    """«token» живёт внутри «tokenizer», а его GigaAM качает каждый раз."""
    from asrhub.engines.gigaam_engine import _load_failure

    отказ = _load_failure("m", [
        "v3_e2e_rnnt: RuntimeError: Download of v3_e2e_rnnt_tokenizer.model "
        "failed after 3 attempts.",
    ], "cpu", "")
    assert "токен" not in str(отказ).lower(), str(отказ)


def test_weights_go_where_the_service_may_write(monkeypatch, tmp_path):
    """Библиотека по умолчанию пишет в ~/.cache/gigaam — службе туда нельзя.

    Юнит работает с ProtectHome=read-only, и загрузка падала с OSError
    [Errno 30] на «/home/asrhub». Переменная GIGAAM_MODEL_DIR, на которую мы
    рассчитывали, не читается никем: библиотека берёт каталог только из
    аргумента download_root.
    """
    import sys
    import types

    from asrhub.catalog import get_model
    from asrhub.engines.gigaam_engine import GigaAMEngine

    вызовы = {}

    поддельный = types.ModuleType("gigaam")

    def load_model(name, device=None, download_root=None):
        вызовы["name"] = name
        вызовы["download_root"] = download_root
        return object()

    поддельный.load_model = load_model
    monkeypatch.setitem(sys.modules, "gigaam", поддельный)

    движок = GigaAMEngine(get_model("gigaam-v3-e2e-rnnt"), {})
    движок._load({"models_dir": str(tmp_path), "device": "cpu"})

    assert вызовы["download_root"] == str(tmp_path), (
        "загрузка снова пойдёт в домашний каталог: " + repr(вызовы))
    assert вызовы["name"] == "v3_e2e_rnnt", вызовы


def test_the_repository_id_is_not_attempted():
    """Идентификатор репозитория библиотека не принимает никогда.

    load_model берёт короткое имя варианта или путь к .ckpt, а на всё прочее
    отвечает перечнем доступных имён. Попытка была не просто бесполезной —
    её ответ и сбивал разбор причины.
    """
    import sys
    import types

    from asrhub.catalog import get_model
    from asrhub.engines.gigaam_engine import GigaAMEngine
    from asrhub.errors import ModelLoadError

    попытки = []
    поддельный = types.ModuleType("gigaam")

    def load_model(name, device=None, download_root=None):
        попытки.append(name)
        raise ValueError(f"Model '{name}' not found. Available model names: ['ssl']")

    поддельный.load_model = load_model
    сохранённый = sys.modules.get("gigaam")
    sys.modules["gigaam"] = поддельный
    try:
        движок = GigaAMEngine(get_model("gigaam-v3-e2e-rnnt"), {})
        with pytest.raises(ModelLoadError):
            движок._load({"models_dir": "", "device": "cpu"})
    finally:
        if сохранённый is None:
            sys.modules.pop("gigaam", None)
        else:
            sys.modules["gigaam"] = сохранённый

    assert попытки == ["v3_e2e_rnnt"], (
        "идентификатор репозитория снова в попытках: " + repr(попытки))


def test_the_first_gigaam_repository_does_not_serve_the_third_version():
    """Голое «rnnt» библиотека сама разворачивает в v3_rnnt.

    Значит репозиторий первой версии молча отдавал третью — подмена, которую
    по выводу не заметить.
    """
    from asrhub.engines.gigaam_engine import variant_name

    for ревизия in ("ctc", "rnnt", "ssl"):
        имя = variant_name("ai-sage/GigaAM", ревизия)
        assert имя.startswith("v1_"), f"{ревизия} → {имя}"


def test_downloaded_gigaam_weights_are_seen_as_installed(tmp_path):
    """Веса GigaAM — файл, а не каталог Hugging Face.

    Поиск умел только раскладку `models--владелец--имя`, и скачанная модель
    показывалась незагруженной навсегда, а её размер — нулевым.
    """
    from asrhub import model_files

    assert model_files.find_local(tmp_path, "ai-sage/GigaAM-v3", "e2e_rnnt") is None
    веса = tmp_path / "v3_e2e_rnnt.ckpt"
    веса.write_bytes(b"x" * 2048)

    найдено = model_files.find_local(tmp_path, "ai-sage/GigaAM-v3", "e2e_rnnt")
    assert найдено == веса, найдено
    assert model_files.directory_size(найдено) == 2048, "размер файла не считается"


def test_a_short_word_cannot_match_inside_a_model_name():
    """Слова разбора должны быть неспособны совпасть с посторонним текстом.

    Правило «разбираем первую попытку» спасает лишь тогда, когда шум пришёл
    из запасных. Если перечень доступных имён напечатала сама первая попытка,
    защищают уже только сами слова: короткое «ssl» совпадёт с именем модели
    «ssl», и отказ «такой модели нет» станет отказом сети.
    """
    from asrhub.engines.gigaam_engine import _load_failure

    отказ = _load_failure("m", [
        "v9_rnnt: ValueError: Model 'v9_rnnt' not found. Available model names: "
        "['ctc', 'rnnt', 'ssl', 'emo', 'v3_ssl']",
    ], "cpu", "")
    assert "хранилищу весов" not in str(отказ), (
        "имя модели в перечне снова выдано за отказ сети: " + str(отказ))


def test_the_results_section_offers_playback(repo_root: Path):
    """Проигрыватель должен быть и в карточке, и строкой списка.

    Проверка по разметке, а не по виду: она ловит случай, когда кнопку
    потеряли при перекраивании списка, — а такое видно только глазами и
    только если открыть нужный раздел.
    """
    app = (repo_root / "server" / "asrhub" / "web" / "app.js").read_text(encoding="utf-8")
    assert 'id="job-player"' in app, "в карточке задания нет проигрывателя"
    assert "__asrhub.playRecording" in app, "в списке результатов нет кнопки прослушивания"
    assert "__asrhub.saveRecording" in app, "запись нельзя скачать"

    # Связь с расшифровкой — то, ради чего всё затевалось.
    assert "segment[data-start]" in app, "щелчок по сегменту не переводит звук"
    assert "classList.add('playing')" in app, "звучащий сегмент не подсвечивается"

    css = (repo_root / "server" / "asrhub" / "web" / "styles.css").read_text(encoding="utf-8")
    assert ".segment.playing" in css, "нет оформления для звучащего сегмента"


def test_the_player_does_not_keep_playing_after_the_card_is_closed(repo_root: Path):
    """Узел удалён, а звук идёт — так ведёт себя <audio>, если его не остановить.

    Проверяем строку целиком, а не соседство двух подстрок в окне на три
    тысячи знаков: окно первым же разрастанием карточки съезжало, и тест
    падал на добавлении соседнего раздела, где всё было в порядке. Здесь же
    остановка и подписка обязаны стоять в одном выражении — иначе это
    подписка неизвестно на что.
    """
    app = (repo_root / "server" / "asrhub" / "web" / "app.js").read_text(encoding="utf-8")
    остановка = [строка for строка in app.splitlines()
                 if "asrhub:closed" in строка and "player.destroy()" in строка]
    assert остановка, "проигрыватель не останавливается вместе с карточкой"


# ---------------------------------------------------------------------------
# Новые разрезы аналитики
# ---------------------------------------------------------------------------


@pytest.fixture()
def rich_db(tmp_path: Path):
    """База с разнообразием, на котором новые разрезы имеют смысл."""
    from asrhub.analytics import Analytics
    from asrhub.db import Database

    database = Database(tmp_path / "rich.sqlite3")
    момент = time.time()
    # Понедельник 10:00 по местному времени — чтобы карта недели была
    # предсказуемой, а не зависела от дня прогона тестов.
    понедельник = момент - (time.localtime(момент).tm_wday * 86400)
    основа = time.mktime(time.localtime(понедельник)[:3] + (10, 0, 0, 0, 0, -1))

    for i in range(6):
        job_id = database.create_job({
            "owner": "alice", "language": "ru", "engine": "gigaam",
            "model": "gigaam-v3-rnnt", "media_duration_s": 120.0,
            "filename": f"разговор-{i}.wav", "file_size": 2_000_000,
            "created_at": основа, "tags": "продажи" if i % 2 else "поддержка",
            "priority": 60 if i < 2 else 50, "source": "phone"})
        database.update_job(job_id, status="completed", finished_at=основа + 40,
                            processing_time_s=20.0, rtf=0.16, queue_time_s=float(i * 5),
                            words_count=300, segments_count=12, avg_confidence=0.9 - i * 0.02,
                            speakers_count=2, device="cuda", peak_memory_mb=2100.0 + i * 10)

    # Задание с повторами, которое так и не дошло.
    сломанное = database.create_job({"owner": "alice", "model": "m", "created_at": основа,
                                     "media_duration_s": 30.0, "filename": "битый.mp3"})
    database.update_job(сломанное, status="failed", retries=2, error_code="decode_error")

    # Задание в другой день недели: без него карту нельзя проверить на то,
    # ради чего она заведена, — что дни не складываются в один.
    среда = database.create_job({"owner": "alice", "model": "m", "media_duration_s": 60.0,
                                 "filename": "среда.wav",
                                 "created_at": основа + 2 * 86400})
    database.update_job(среда, status="completed", finished_at=основа + 2 * 86400 + 30,
                        processing_time_s=10.0, queue_time_s=1.0, device="cpu")

    # Повтор уже виденного файла.
    повтор = database.create_job({"owner": "alice", "model": "gigaam-v3-rnnt",
                                  "created_at": основа, "media_duration_s": 120.0,
                                  "filename": "разговор-0.wav", "file_hash": "abc"})
    database.update_job(повтор, status="completed", cached_from="job_первое",
                        finished_at=основа + 1, queue_time_s=0.2)
    return Analytics(database)


def test_the_week_map_separates_weekdays_from_weekends(rich_db):
    """Суточный профиль складывает понедельник с воскресеньем в одно число.

    Планировать по нему обслуживание нельзя: у телефонии утро понедельника и
    вечер воскресенья — разные миры, а в одной строке они неразличимы.
    """
    карта = rich_db.weekly_heatmap("month")
    assert len(карта["jobs"]) == 7 and len(карта["jobs"][0]) == 24
    assert карта["total"] == 9, карта["total"]

    # Главное: задания в разные дни лежат в разных строках. Если дни
    # схлопнуть, непустой строкой останется одна — и карта превратится в тот
    # же суточный профиль, только выше.
    непустые = [i for i, строка in enumerate(карта["jobs"]) if sum(строка)]
    assert len(непустые) == 2, f"дни сложились в один: {непустые}"
    assert непустые[1] - непустые[0] == 2, непустые

    assert карта["peak"]["jobs"] == 8, карта["peak"]
    assert карта["peak"]["hour"] == 10, карта["peak"]
    сумма = sum(v for строка in карта["jobs"] for v in строка)
    assert сумма == карта["total"], "сумма по карте разошлась с итогом"


def test_reliability_separates_the_first_attempt_from_the_third(rich_db):
    """Доля успеха не отличает безупречное задание от прошедшего с третьего раза."""
    r = rich_db.reliability("month")
    assert r["total"] == 9
    assert r["jobs_with_retries"] == 1, r
    assert r["retry_total"] == 2, r
    # Повторявшееся задание не дошло — значит со второй попытки не дошёл никто.
    assert r["completed_after_retry"] == 0, r
    assert r["first_attempt_success"] == 8, r


def test_cache_savings_are_counted_in_hours_not_in_hits(rich_db):
    """«Из кеша: 1» не отвечает на вопрос, стоило ли оно того."""
    c = rich_db.cache_savings("month")
    assert c["hits"] == 1, c
    # Две минуты записи, снятые повтором.
    assert c["audio_hours_saved"] == pytest.approx(120 / 3600, rel=0.01), c
    assert c["assumed_rtf"] and c["assumed_rtf"] > 0, "не по чему считать экономию"
    assert c["processing_seconds_saved"] > 0, c


def test_queue_latency_shows_the_tail_not_the_average(rich_db):
    """На среднее ожидание не жалуются — жалуются на хвост."""
    q = rich_db.queue_latency("month")
    assert q["overall"]["count"] == 8, q["overall"]
    assert q["overall"]["p95"] >= q["overall"]["p50"], "перцентили не по порядку"
    группы = {r["name"] for r in q["by_priority"]}
    assert "высокий (>50)" in группы and "обычный (50)" in группы, группы


def test_audio_profile_describes_the_material(rich_db):
    """Разрез не про сервер, а про то, что на него приносят."""
    a = rich_db.audio_profile("month")
    форматы = {f["format"]: f for f in a["formats"]}
    assert "wav" in форматы, форматы
    # 300 слов за две минуты — 150 слов в минуту.
    assert a["speech_rate_wpm"]["avg"] == pytest.approx(150, rel=0.02), a["speech_rate_wpm"]
    assert a["speakers"] and a["speakers"][0]["speakers"] == 2, a["speakers"]


def test_resources_answer_the_question_asked_before_buying_a_card(rich_db):
    """Сколько памяти просит модель на пике — и что уехало на процессор."""
    r = rich_db.resources("month")
    модели = {m["model"]: m for m in r["models"]}
    assert "gigaam-v3-rnnt" in модели, модели
    assert модели["gigaam-v3-rnnt"]["peak_mb"] == pytest.approx(2150, rel=0.01), модели
    устройства = {d["device"]: d for d in r["devices"]}
    assert "cuda" in устройства, устройства


def test_tags_are_the_only_breakdown_the_user_defines(rich_db):
    """Модель и движок сервер знает про себя, а метка отвечает «на какой проект»."""
    метки = {t["tag"]: t for t in rich_db.by_tag("month")}
    assert "продажи" in метки and "поддержка" in метки, метки
    assert метки["продажи"]["jobs"] == 3, метки["продажи"]
    assert метки["продажи"]["audio_hours"] > 0, метки["продажи"]


def test_quality_trend_has_a_point_where_there_is_data(rich_db):
    """Средняя уверенность за месяц — число ни о чём; полезен ход."""
    q = rich_db.quality_trend("month", buckets=12)
    assert len(q["buckets"]) == 12
    заполнено = [v for v in q["confidence"] if v is not None]
    assert заполнено, "ход уверенности пуст при наличии данных"
    assert 0 < заполнено[-1] <= 1, заполнено


def test_every_new_section_is_reachable_by_its_own_url(repo_root: Path):
    """Разделы аналитики забирают по одному — сводный отчёт весит десятки килобайт."""
    источник = (repo_root / "server" / "asrhub" / "api" / "routes_system.py").read_text(
        encoding="utf-8")
    for раздел in ("weekly", "cache", "reliability", "audio", "resources",
                   "quality", "tags", "queue"):
        assert f'"{раздел}": state.analytics.' in источник, f"раздел {раздел} не отдаётся"


def test_peak_memory_counter_is_reset_before_the_job_not_after(monkeypatch):
    """Счётчик пика у torch общий на процесс и копится с самого запуска.

    Если обнулять его после замера, первое задание отчитается за всё, что
    успело выделиться до него. Сброс должен быть в начале — и до того, как
    движок что-то посчитает.
    """
    from asrhub import processor as proc

    события: list[str] = []

    class ФейковаяCuda:
        @staticmethod
        def is_available() -> bool:
            return True

        @staticmethod
        def reset_peak_memory_stats() -> None:
            события.append("сброс")

        @staticmethod
        def max_memory_allocated() -> int:
            события.append("замер")
            return 700 * 1024 * 1024

    класс_torch = type(sys)("torch")
    класс_torch.cuda = ФейковаяCuda
    monkeypatch.setitem(sys.modules, "torch", класс_torch)

    proc._reset_peak_memory()
    assert proc._peak_memory_mb("cuda:0") == 700.0
    assert события == ["сброс", "замер"], события


def test_process_job_resets_the_counter_only_when_it_is_alone(tmp_path: Path,
                                                             monkeypatch):
    """Сброс — не бесплатная операция для соседа по очереди.

    Замер делается всегда: иначе на занятом сервере — то есть там, где
    вопрос о памяти и стоит, — раздел «Ресурсы» пуст. А вот обнулять общий
    счётчик, когда рядом идёт второе задание, нельзя: это стирает то, что
    сосед уже накопил, и его собственный замер выходит заниженным.
    """
    from asrhub import processor as proc
    from asrhub.errors import AudioError

    сбросов: list[int] = []
    monkeypatch.setattr(proc, "_reset_peak_memory", lambda: сбросов.append(1))

    # Задание падает сразу на несуществующем файле — сброс к тому моменту
    # уже должен был случиться: он идёт до всякой работы.
    источник = tmp_path / "нет.wav"
    for обнулять, ожидание in ((False, 0), (True, 1)):
        сбросов.clear()
        with pytest.raises(AudioError):
            proc.process_job(источник, {"model": "нет-такой"},
                             registry=None, workdir=tmp_path, outdir=tmp_path,
                             basename="x", measure_memory=True,
                             reset_memory=обнулять)
        assert len(сбросов) == ожидание, (обнулять, сбросов)


def test_gpu_samples_are_cleaned_like_every_other_metric(tmp_path: Path):
    """Замеры по картам — единственная таблица, которую не чистил никто.

    Пишется по строке на карту за такт, то есть на машине с двумя картами
    растёт вдвое быстрее общих замеров. Срок хранения у обеих таблиц один.
    """
    from asrhub.db import Database

    database = Database(tmp_path / "g.sqlite3")
    сейчас = time.time()
    давно = сейчас - 400 * 86400
    for метка in (давно, сейчас):
        database.add_system_sample({"ts": метка, "cpu_percent": 10.0})
        database.add_gpu_samples(метка, [
            {"gpu": 0, "name": "RTX 5090", "util_percent": 50.0,
             "mem_used_mb": 1000.0},
            {"gpu": 1, "name": "RTX 5090", "util_percent": 60.0,
             "mem_used_mb": 1100.0},
        ])

    removed = database.cleanup(results_days=30, metrics_days=30)
    assert removed["gpu_samples"] == 2, removed
    осталось = database.gpu_samples(since=0)
    assert len(осталось) == 2, осталось
    assert all(r["ts"] >= сейчас - 1 for r in осталось), осталось


def test_the_report_reads_the_archive_once_not_twenty_times(rich_db):
    """Два десятка разрезов — два десятка полных проходов по архиву.

    Каждый разрез читал базу сам: на архиве в сорок тысяч заданий сводный
    отчёт собирался тринадцать секунд вместо двух. Разрезы по завершённым и
    по упавшим — отдельные выборки, поэтому проходов остаётся несколько, но
    не по одному на разрез.
    """
    запросы: list[tuple[Any, ...]] = []
    исходный = rich_db.db.list_jobs

    def учёт(**kw):
        запросы.append((kw.get("since"), kw.get("status"), kw.get("owner")))
        return исходный(**kw)

    rich_db.db.list_jobs = учёт
    try:
        отчёт = rich_db.full_report("month")
    finally:
        rich_db.db.list_jobs = исходный

    assert len(отчёт) >= 19, "разрезы потерялись"
    assert len(запросы) <= 5, f"{len(запросы)} проходов по архиву: {запросы}"
    assert len(set(запросы)) == len(запросы), "один и тот же запрос выполнен дважды"


def test_all_sections_of_one_report_cover_the_same_window(rich_db):
    """Иначе «месяц» у первого разреза начинался раньше, чем у последнего.

    Границу окна каждый разрез считал сам, от текущего времени. За секунды
    сборки отчёта она уезжала — разрезы расходились между собой, и общая
    выборка не попадала в кеш ни разу.
    """
    границы: list[float] = []
    исходный = rich_db.db.list_jobs

    def учёт(**kw):
        if kw.get("since") is not None:
            границы.append(float(kw["since"]))
        return исходный(**kw)

    rich_db.db.list_jobs = учёт
    try:
        rich_db.full_report("month")
    finally:
        rich_db.db.list_jobs = исходный

    assert границы, "ни один разрез не ограничил окно"
    assert len(set(границы)) == 1, f"окна разъехались: {sorted(set(границы))}"


def test_outside_a_report_each_section_still_reads_fresh(rich_db):
    """Кеш живёт ровно один отчёт — иначе это просто устаревшие данные.

    Аналитика существует всё время работы приложения; выборка, пережившая
    свой отчёт, показывала бы вчерашний архив как сегодняшний.
    """
    assert rich_db._выборки is None, "кеш остался включённым после отчёта"
    rich_db.overview("month")
    assert rich_db._выборки is None, "разрез вне отчёта включил кеш"


def test_a_recording_with_an_ascii_name_keeps_it_on_save(client, tmp_path: Path):
    """Кнопка «Скачать запись» теряла имя у обычных латинских файлов.

    Заголовок сервер шлёт по-разному: `filename*=` по RFC 5987 появляется
    только у имён с кириллицей, а у «record.wav» его нет. Разбор в плеере
    искал только его, и такой файл сохранялся как «запись-job_….wav».
    """
    исходник = tmp_path / "record.wav"
    _wav_with_lead_silence(исходник, silence_s=0.1, total_s=0.5)

    with исходник.open("rb") as fh:
        ответ = client.post("/api/jobs", files={"file": ("record.wav", fh, "audio/wav")})
    assert ответ.status_code in (200, 201), ответ.text
    job_id = ответ.json()["id"]

    аудио = client.get(f"/api/jobs/{job_id}/audio")
    assert аудио.status_code == 200, аудио.text
    заголовок = аудио.headers.get("content-disposition", "")
    assert "record.wav" in заголовок, заголовок
    assert "filename*=" not in заголовок, (
        "если сервер стал слать RFC 5987 и для латиницы — проверка устарела")

    # Плеер разбирает заголовок общей функцией, а не своей копией.
    app_js = (Path(__file__).resolve().parent.parent / "server" / "asrhub"
              / "web" / "app.js").read_text(encoding="utf-8")
    сохранение = app_js[app_js.index("  async save(id) {"):]
    сохранение = сохранение[:сохранение.index("\n  },")]
    assert "parseFilename(disposition)" in сохранение, сохранение
    assert "filename\\*=utf-8" not in сохранение, "своя копия разбора вернулась"


def test_a_status_code_is_not_recognised_inside_a_file_path():
    """«401» в имени файла отправляло чинить токен вместо загрузки весов.

    Приметы кодов состояния были обычными подстроками, и «401» находилось
    внутри «v3_rnnt-8401.ckpt» и внутри номера порта прокси. Отсутствующий
    на диске файл сервер объявлял отказом в доступе к Hugging Face.
    """
    from asrhub.engines.gigaam_engine import _load_failure

    def причина(*ошибки: str) -> str:
        текст = str(_load_failure("gigaam-v3-rnnt", list(ошибки), "cuda", "/opt"))
        return текст.split("«gigaam-v3-rnnt»", 1)[1].split(".")[0].strip(": ")

    assert причина(
        "v3_rnnt: FileNotFoundError: [Errno 2] No such file or directory: "
        "'/opt/asrhub/models/gigaam/v3_rnnt-8401.ckpt'"
    ) == "веса не найдены на диске"

    assert причина(
        "ConnectionError: HTTPConnectionPool(host='proxy', port=8403): "
        "Max retries exceeded"
    ) == "сервер не смог обратиться к хранилищу весов"

    # А настоящие коды по-прежнему опознаются.
    assert причина(
        "HTTPError: 401 Client Error: Unauthorized for url: https://huggingface.co/x"
    ) == "к весам нужен доступ по токену Hugging Face"
    assert причина(
        "GatedRepoError: 403 Client Error. Access to model is restricted."
    ) == "к весам нужен доступ по токену Hugging Face"


def test_the_two_cache_savings_figures_agree(rich_db):
    """На одной странице стояли два числа про экономию от кеша.

    Разрез «эффективность» считал её по вбитой в код скорости 0.2, а
    разрез «кеш» — по измеренной на своих же заданиях. Расходились вдвое.
    """
    отчёт = rich_db.full_report("month")
    э, к = отчёт["efficiency"], отчёт["cache"]

    assert э["assumed_rtf"] == к["assumed_rtf"], (э["assumed_rtf"], к["assumed_rtf"])
    assert э["assumed_rtf"], "скорость не измерена — экономию не по чему считать"
    assert э["cache_hits"] == к["hits"], (э["cache_hits"], к["hits"])
    # Часы округлены до трёх знаков — это 3.6 секунды; сравниваем в пределах
    # половины этого шага, иначе тест меряет округление, а не согласие.
    assert э["saved_compute_hours"] == pytest.approx(
        к["processing_seconds_saved"] / 3600, abs=0.0005), (э, к)


def test_the_measured_speed_ignores_jobs_taken_from_cache(rich_db):
    """У задания из кеша своего времени обработки нет.

    Включать его в среднее — занижать цену работы тем сильнее, чем чаще
    срабатывает кеш, то есть тем сильнее, чем важнее ответ.
    """
    задания = [
        {"status": "completed", "media_duration_s": 100.0, "processing_time_s": 50.0},
        {"status": "completed", "media_duration_s": 100.0, "cached_from": "job_x",
         "processing_time_s": 0.0},
        {"status": "failed", "media_duration_s": 100.0, "processing_time_s": 90.0},
    ]
    assert rich_db._measured_rtf(задания) == pytest.approx(0.5), "кеш попал в среднее"


def test_every_accepted_format_has_a_content_type_for_playback():
    """Ответ, который браузер не станет играть, — тот же отказ, без объяснения.

    В таблице было двенадцать расширений из двадцати девяти принимаемых, а
    остальные уходили в `mimetypes.guess_type`, который про «.caf», «.w64» и
    «.m2ts» не знает ничего: браузер получал «application/octet-stream».
    """
    from asrhub.api.routes_jobs import _AUDIO_TYPES
    from asrhub.pipeline.audio import SUPPORTED_EXTENSIONS

    без_типа = sorted(SUPPORTED_EXTENSIONS - set(_AUDIO_TYPES))
    assert not без_типа, f"принимаем, но не отдаём на прослушивание: {без_типа}"
    лишние = sorted(set(_AUDIO_TYPES) - SUPPORTED_EXTENSIONS)
    assert not лишние, f"тип есть, а файл такой сервер не принимает: {лишние}"
    for расш, тип in _AUDIO_TYPES.items():
        assert тип.startswith(("audio/", "video/")), (расш, тип)


def test_an_unknown_gigaam_repository_is_not_silently_taken_for_v3(caplog):
    """Голое «rnnt» библиотека разворачивает в v3_rnnt — молча.

    Если в справочник добавят четвёртую версию, не поправив таблицу
    вариантов, сервер будет уверенно качать и грузить веса третьей. Угадать
    тут нечего, но сказать об этом надо.
    """
    from asrhub.engines.gigaam_engine import weights_file

    with caplog.at_level("WARNING"):
        assert weights_file("ai-sage/GigaAM-v4", "rnnt") == "v3_rnnt.ckpt"
    assert any("не описан в таблице вариантов" in r.message for r in caplog.records), \
        caplog.records

    # Известные репозитории предупреждений не дают.
    caplog.clear()
    with caplog.at_level("WARNING"):
        assert weights_file("ai-sage/GigaAM", "rnnt") == "v1_rnnt.ckpt"
        assert weights_file("ai-sage/GigaAM-v3", "e2e_rnnt") == "v3_e2e_rnnt.ckpt"
    assert not caplog.records, caplog.records


def test_the_quality_axis_is_labelled_by_the_width_of_its_bucket(rich_db, repo_root):
    """На часовом окне ход качества подписывал все точки одной датой.

    Подпись была жёстко «день.месяц», а корзина за час — две с половиной
    минуты: двадцать четыре одинаковых подписи вместо оси.
    """
    # Своя свежая запись: общая заготовка кладёт задания на понедельник
    # текущей недели, и часовое окно попадало на них ровно один час в
    # неделю. Тест из-за этого проходил по понедельникам и падал в
    # остальные дни — про часовую подпись он при этом не проверял ничего.
    недавнее = rich_db.db.create_job({
        "owner": "alice", "model": "gigaam-v3-rnnt", "engine": "gigaam",
        "language": "ru", "media_duration_s": 90.0, "filename": "свежая.wav",
        "created_at": time.time() - 600})
    rich_db.db.update_job(недавнее, status="completed", finished_at=time.time() - 500,
                          processing_time_s=15.0, rtf=0.17, words_count=200,
                          segments_count=8, avg_confidence=0.88)

    for период, шире in (("hour", 7200), ("month", 0)):
        ход = rich_db.quality_trend(период)
        assert "bucket_seconds" in ход, "ширина корзины не отдаётся"
        assert ход["bucket_seconds"] > 0, ход["bucket_seconds"]
        if период == "hour":
            assert ход["bucket_seconds"] < шире, (
                "часовое окно должно подписываться временем, а не датой")

    # Пустой ход тоже отдаёт поле — иначе подпись выбирается наугад.
    пустой = rich_db.quality_trend("hour", buckets=4)
    assert "bucket_seconds" in пустой

    app_js = (repo_root / "server" / "asrhub" / "web" / "app.js").read_text(
        encoding="utf-8")
    assert "подписьВремени(t, qt.bucket_seconds)" in app_js, \
        "ход качества снова подписывается сам по себе"
    assert "подписьВремени(t, ts.bucket_seconds)" in app_js, \
        "поток заданий перестал пользоваться общей подписью"


def test_the_quality_cards_say_something_when_there_is_nothing(repo_root):
    """Две дырки без объяснения — худший из возможных ответов.

    Когда завершённых заданий за период нет, разрез отдаёт пустые корзины, и
    карточки оставались нарисованными, но пустыми внутри.
    """
    app_js = (repo_root / "server" / "asrhub" / "web" / "app.js").read_text(
        encoding="utf-8")
    начало = app_js.index("const qt = data.quality_trend")
    кусок = app_js[начало:начало + 900]
    assert "!(qt.buckets || []).length" in кусок, кусок[:300]
    assert кусок.count("Charts.empty") >= 2, "пустые карточки снова молчат"


def test_the_interface_is_not_served_from_a_stale_browser_cache(client, repo_root: Path):
    """После обновления сервера в браузере оставался прежний интерфейс.

    Starlette не ставит на статику `Cache-Control` вовсе, а ответ без явного
    срока браузер волен держать по своему усмотрению — обычно десятую часть
    возраста файла. Для двухмесячного `app.js` это почти неделя, в течение
    которой браузер не спрашивает сервер ни разу: обновление проходило,
    файлы на диске менялись, а человек видел старую версию.

    Воспроизведено в настоящем Chromium: файл с давней отметкой времени,
    подмена содержимого, обычная перезагрузка — и прежний скрипт из кеша.
    """
    страница = client.get("/")
    assert страница.status_code == 200
    assert страница.headers.get("cache-control") == "no-cache", (
        "саму страницу кешировать нельзя: в ней лежат ссылки на всё остальное")

    ссылки = re.findall(r"/static/([A-Za-z0-9_.-]+)\?v=([a-f0-9]+)", страница.text)
    имена = {имя for имя, _ in ссылки}
    assert {"app.js", "charts.js", "styles.css"} <= имена, имена
    assert all(len(отпечаток) >= 8 for _, отпечаток in ссылки), ссылки

    # Отпечаток — по содержимому: правка файла обязана менять адрес.
    app_js = repo_root / "server" / "asrhub" / "web" / "app.js"
    было = app_js.read_text(encoding="utf-8")
    прежний = dict(ссылки)["app.js"]
    try:
        app_js.write_text(было + "\n// проверка отпечатка\n", encoding="utf-8")
        новый = dict(re.findall(r"/static/([A-Za-z0-9_.-]+)\?v=([a-f0-9]+)",
                                client.get("/").text))["app.js"]
    finally:
        app_js.write_text(было, encoding="utf-8")
    assert новый != прежний, "правка файла не изменила адрес — обновление снова не дойдёт"

    # Со отпечатком содержимое неизменно, поэтому его можно держать долго;
    # без отпечатка за свежесть отвечать нечему — только с перепроверкой.
    с_меткой = client.get(f"/static/app.js?v={прежний}")
    assert "immutable" in с_меткой.headers.get("cache-control", ""), с_меткой.headers
    без_метки = client.get("/static/app.js")
    assert без_метки.headers.get("cache-control") == "no-cache", без_метки.headers


# ---------------------------------------------------------------------------
# Поиск по расшифровкам
# ---------------------------------------------------------------------------

def _архив(path: Path, разговоры: list[tuple[str, list[str]]]):
    """База с расшифровками: имя файла и реплики."""
    from asrhub.db import Database

    database = Database(path)
    номера = []
    for имя, реплики in разговоры:
        job_id = database.create_job({"model": "gigaam-v3-rnnt", "filename": имя})
        database.update_job(job_id, status="completed", text=" ".join(реплики))
        database.save_segments(job_id, [
            {"start": i * 10.0, "end": i * 10.0 + 9.0, "text": t,
             "speaker": f"S{i % 2}"}
            for i, t in enumerate(реплики)])
        номера.append(job_id)
    return database, номера


def test_search_finds_the_phrase_and_says_where_it_was_said(tmp_path: Path):
    """Список отвечал «нашлось в этом разговоре» и замолкал.

    Дальше человек открывал карточку и искал глазами — при том что сервер
    уже знал и фразу, и секунду, на которой она сказана.
    """
    database, (первый, второй) = _архив(tmp_path / "a.sqlite3", [
        ("переговоры.wav", ["Добрый день, обсудим договор поставки",
                            "Сроки нас не устраивают",
                            "Клиент грозит передать спор в арбитраж"]),
        ("поддержка.wav", ["Не работает личный кабинет",
                           "Попробуйте сбросить пароль"]),
    ])
    assert database.fts_ready, "указатель не собрался — остальное проверять нечего"

    найдено = database.list_jobs(search="арбитраж", limit=10, light=True)
    assert [j["id"] for j in найдено] == [первый], найдено

    реплики = database.search_segments("арбитраж")
    assert len(реплики) == 1, реплики
    r = реплики[0]
    assert r["job_id"] == первый
    assert r["start_s"] == 20.0, r
    assert "‹арбитраж›" in r["snippet"], r["snippet"]

    # Ищется и по началу слова: поиск должен работать по ходу набора.
    assert database.jobs_matching("арбитр") == [первый]
    # И без разницы «ещё»/«еще» — иначе надо угадывать, как набрано.
    assert database.jobs_matching("несуществующее") == []
    assert второй not in database.jobs_matching("договор")


def test_the_index_forgets_what_was_deleted(tmp_path: Path):
    """Реплики удаляются из четырёх мест, и любое забывшее оставит призраков.

    Поэтому указатель держат триггеры, а не код: код можно забыть поправить,
    триггер лежит в самой схеме.
    """
    database, (первый, второй) = _архив(tmp_path / "b.sqlite3", [
        ("один.wav", ["говорим про арбитраж"]),
        ("два.wav", ["говорим про доставку"]),
    ])
    assert database.jobs_matching("арбитраж") == [первый]

    database.delete_job(первый)
    assert database.jobs_matching("арбитраж") == [], "удалённый разговор ищется"
    assert database.jobs_matching("доставку") == [второй], "заодно потерялся живой"


def test_re_running_a_job_does_not_leave_the_old_words_in_the_index(tmp_path: Path):
    """Повтор задания заменяет реплики целиком.

    Если старые остаются в указателе, поиск находит слова, которых в
    расшифровке уже нет, — и открытая карточка их не содержит.
    """
    database, (job_id,) = _архив(tmp_path / "c.sqlite3", [
        ("запись.wav", ["прежнее слово арбитраж"]),
    ])
    assert database.jobs_matching("арбитраж") == [job_id]

    database.save_segments(job_id, [
        {"start": 0.0, "end": 5.0, "text": "новое слово доставка"}])
    assert database.jobs_matching("арбитраж") == [], "старое слово всё ещё ищется"
    assert database.jobs_matching("доставка") == [job_id]


def test_search_still_works_without_the_index(tmp_path: Path):
    """FTS5 — необязательный модуль SQLite.

    На сборке без него сервер обязан подняться и искать перебором: медленно,
    но искать. Сервер, который не стартует, не работает никак.
    """
    database, (первый, _) = _архив(tmp_path / "d.sqlite3", [
        ("один.wav", ["говорим про арбитраж"]),
        ("два.wav", ["говорим про доставку"]),
    ])
    database.fts_ready = False
    найдено = database.list_jobs(search="арбитраж", limit=10, light=True)
    assert [j["id"] for j in найдено] == [первый], найдено
    # Фразы без указателя нет — и это честнее, чем выдумать её.
    assert database.search_segments("арбитраж") == []


def test_a_query_typed_by_a_person_is_never_a_syntax_error(tmp_path: Path):
    """У FTS5 свой язык запросов.

    Одинокая кавычка, звёздочка или слово AND — это не поиск, а
    синтаксическая ошибка прямо в лицо человеку, который искал «договор».
    """
    database, (job_id,) = _архив(tmp_path / "e.sqlite3", [
        ("один.wav", ["обсудили договор и сроки"]),
    ])
    for запрос in ('"', '*', 'AND', 'OR NOT', 'договор"', '(', ')', '^', '-',
                   'NEAR(', 'договор AND', '  ', 'a' * 300, 'договор ' * 40):
        найдено = database.list_jobs(search=запрос, limit=5, light=True)
        assert isinstance(найдено, list), запрос

    assert database.list_jobs(search="договор", limit=5, light=True), "обычный запрос сломался"


def test_an_existing_archive_becomes_searchable_on_upgrade(tmp_path: Path):
    """У накопленного архива указателя ещё нет.

    Без наполнения при первом открытии поиск не нашёл бы ни одного старого
    разговора — то есть возможность, ради которой всё делалось, не работала
    бы ровно там, где она нужнее всего.
    """
    from asrhub.db import Database

    путь = tmp_path / "f.sqlite3"
    database, (job_id,) = _архив(путь, [("старая.wav", ["древнее слово арбитраж"])])
    # Ровно то состояние, в котором база приезжает с прошлой версии:
    # реплики есть, указателя нет.
    database.execute("DROP TABLE IF EXISTS segments_fts")
    database.execute("PRAGMA user_version=7")
    database.close()

    заново = Database(путь)
    assert заново.fts_ready, "указатель не завёлся на существующей базе"
    assert заново.jobs_matching("арбитраж") == [job_id], "старый архив не проиндексирован"


# ---------------------------------------------------------------------------
# Действия над выборкой
# ---------------------------------------------------------------------------

def test_one_bad_job_does_not_cancel_the_whole_batch(client, tmp_path: Path):
    """В выборку почти всегда попадает что-то, к чему действие неприменимо.

    Прерывать всю команду из-за одной такой строки означало бы, что
    пакетное действие работает только на идеально подобранной выборке — то
    есть почти никогда.
    """
    номера = []
    for i in range(3):
        файл = tmp_path / f"з-{i}.wav"
        _wav_with_lead_silence(файл, silence_s=0.05, total_s=0.4)
        with файл.open("rb") as fh:
            ответ = client.post("/api/jobs",
                                files={"file": (файл.name, fh, "audio/wav")})
        assert ответ.status_code in (200, 201), ответ.text
        номера.append(ответ.json()["id"])

    ответ = client.post("/api/jobs/bulk", json={
        "action": "tag", "ids": [*номера, "job_несуществующее"], "tags": "продажи"})
    assert ответ.status_code == 200, ответ.text
    итог = ответ.json()
    assert sorted(итог["done"]) == sorted(номера), итог
    assert [f["id"] for f in итог["failed"]] == ["job_несуществующее"], итог
    # Причина — строка для человека, а не словарь в кавычках.
    assert "не найдено" in итог["failed"][0]["error"], итог["failed"][0]

    for job_id in номера:
        assert client.get(f"/api/jobs/{job_id}").json()["tags"] == "продажи"


def test_a_batch_refuses_an_unknown_action_and_says_which_it_knows(client):
    """Молчаливое «ничего не произошло» на опечатке — худший ответ."""
    ответ = client.post("/api/jobs/bulk",
                        json={"action": "взорвать", "ids": ["job_x"]})
    assert ответ.status_code == 400, ответ.text
    тело = ответ.json()
    assert "взорвать" in тело["message"], тело
    for действие in ("retry", "delete", "tag"):
        assert действие in тело["hint"], тело["hint"]


def test_a_batch_has_a_ceiling(client):
    """Пакет в десятки тысяч заданий занял бы блокировку записи на минуты.

    Очередь на это время встала бы целиком: каждое действие — это запись в
    базу и работа с файлами.
    """
    from asrhub.api.routes_jobs import BULK_LIMIT

    ответ = client.post("/api/jobs/bulk", json={
        "action": "tag", "ids": [f"job_{i}" for i in range(BULK_LIMIT + 1)]})
    assert ответ.status_code == 400, ответ.text
    assert str(BULK_LIMIT) in ответ.json()["message"], ответ.json()

    ответ = client.post("/api/jobs/bulk", json={"action": "tag", "ids": []})
    assert ответ.status_code == 400, ответ.text


def test_a_batch_does_not_touch_other_owners_jobs(auth_client, two_users):
    """Выборку присылает клиент, и в ней может оказаться что угодно.

    Пакетное действие обязано проверять права на каждое задание отдельно —
    иначе оно становится способом удалить чужой архив, зная только номера.
    """
    чужой_id = two_users["job_id"]                 # задание Алисы

    ответ = auth_client.post("/api/jobs/bulk",
                             json={"action": "delete", "ids": [чужой_id]},
                             headers={"X-API-Key": two_users["bob"]})
    assert ответ.status_code == 200, ответ.text
    итог = ответ.json()
    assert итог["done"] == [], "пакет удалил чужое задание"
    assert [f["id"] for f in итог["failed"]] == [чужой_id], итог

    # Задание на месте — видно администратору.
    цел = auth_client.get(f"/api/jobs/{чужой_id}",
                          headers={"X-API-Key": two_users["admin"]})
    assert цел.status_code == 200, цел.text


def test_deleting_one_and_deleting_many_do_the_same_thing(repo_root: Path):
    """Две копии удаления разъедутся: одна забудет файлы, вторая — записи.

    Поэтому обе ручки зовут одну функцию.
    """
    текст = (repo_root / "server" / "asrhub" / "api" / "routes_jobs.py").read_text(
        encoding="utf-8")
    assert текст.count("def _delete_one(") == 1, "копий удаления стало больше одной"
    начало = текст.index('@router.delete("/{job_id}"')
    одиночное = текст[начало:начало + 700]
    assert "_delete_one(" in одиночное, "одиночное удаление снова живёт своей жизнью"
    assert "shutil.rmtree" not in одиночное, одиночное


# ---------------------------------------------------------------------------
# Пик памяти на занятом сервере
# ---------------------------------------------------------------------------

def test_peak_memory_is_measured_even_when_the_server_is_busy(repo_root: Path):
    """Замер делался только когда задание в очереди одно.

    То есть на сервере с двумя воркерами — там, где вопрос «хватит ли
    карты» и стоит, — раздел «Ресурсы» оставался пустым всегда.
    """
    очередь = (repo_root / "server" / "asrhub" / "job_queue.py").read_text(
        encoding="utf-8")
    assert "measure_memory=True" in очередь, "замер снова стал условным"
    assert "reset_memory=одно" in очередь, (
        "обнуление должно оставаться условным: иначе оно стирает то, "
        "что накопил сосед по очереди")


def test_the_peak_is_recorded_together_with_how_many_ran_at_once(tmp_path: Path):
    """«27 ГБ» — это ответ или нет, смотря сколько заданий шло разом.

    Счётчики памяти общие на процесс и по модели её не делят; число
    одновременных заданий превращает бесполезную цифру в ответ на вопрос
    «сколько их выдержит карта».
    """
    from asrhub.analytics import Analytics
    from asrhub.db import Database

    database = Database(tmp_path / "m.sqlite3")
    момент = time.time()
    for разом, память in ((1, 9000.0), (1, 9200.0), (2, 17500.0), (3, 26000.0)):
        job_id = database.create_job({"model": "gigaam-v3-rnnt",
                                      "media_duration_s": 60.0,
                                      "created_at": момент})
        database.update_job(job_id, status="completed", finished_at=момент + 10,
                            processing_time_s=10.0, device="cuda",
                            peak_memory_mb=память, peak_memory_jobs=разом)

    разрез = Analytics(database).resources("month")
    строки = {c["jobs_at_once"]: c for c in разрез["concurrency"]}
    assert set(строки) == {1, 2, 3}, разрез["concurrency"]
    assert строки[1]["measurements"] == 2, строки[1]
    assert строки[3]["peak_mb"] == 26000.0, строки[3]
    # Разрез по моделям при этом остаётся: он отвечает на другой вопрос.
    assert разрез["models"][0]["model"] == "gigaam-v3-rnnt", разрез["models"]


def test_concurrency_counts_the_spike_not_the_ends(tmp_path: Path, monkeypatch):
    """Всплеск в середине — ровно тот случай, когда памяти и не хватает.

    Считать одновременность по началу и концу задания значит его не
    заметить.
    """
    from asrhub.job_queue import JobQueue

    очередь = JobQueue.__new__(JobQueue)
    очередь._running = {}
    очередь._concurrency = {}

    очередь._running["a"] = 0.0
    очередь._concurrency["a"] = 0
    очередь._note_concurrency()
    assert очередь._concurrency["a"] == 1

    # Всплеск: пока «a» работает, приходят и уходят двое.
    for имя in ("b", "c"):
        очередь._running[имя] = 0.0
        очередь._concurrency[имя] = 0
        очередь._note_concurrency()
    for имя in ("b", "c"):
        очередь._running.pop(имя)
        очередь._concurrency.pop(имя)
        очередь._note_concurrency()

    assert len(очередь._running) == 1, "состояние разъехалось"
    assert очередь._concurrency["a"] == 3, (
        f"всплеск не замечен: {очередь._concurrency['a']}")


# ---------------------------------------------------------------------------
# Выгрузка аналитики
# ---------------------------------------------------------------------------

def test_the_report_can_be_taken_away_as_a_table(client):
    """Отчёт можно было только смотреть.

    Чтобы отдать месячные числа руководителю, их переписывали руками — и
    переписывали с округлённых значений на экране, а не с тех, что посчитал
    сервер.
    """
    import openpyxl

    ответ = client.get("/api/analytics/export?period=month&fmt=xlsx")
    assert ответ.status_code == 200, ответ.text
    assert "spreadsheetml" in ответ.headers["content-type"], ответ.headers
    assert "xlsx" in ответ.headers["content-disposition"], ответ.headers
    # Выгрузка считается на момент запроса: закешированная — это вчерашние
    # числа под сегодняшним именем.
    assert ответ.headers.get("cache-control") == "no-store", ответ.headers

    книга = openpyxl.load_workbook(io.BytesIO(ответ.content))
    assert книга.sheetnames, "книга без листов"
    # Подписи русские: файл уходит бухгалтеру, а не разработчику.
    for лист in книга.worksheets:
        заголовки = [c.value for c in лист[1]]
        assert all(isinstance(з, str) and з for з in заголовки), (лист.title, заголовки)
        assert any(re.search(r"[А-Яа-яЁё]", str(з)) for з in заголовки), \
            f"лист «{лист.title}» подписан не по-русски: {заголовки}"
        assert лист.freeze_panes == "A2", лист.title


def test_the_csv_archive_opens_in_excel_by_double_click(client):
    """«Правильный» CSV с запятыми Excel с русскими настройками не разбирает.

    Он показывает его одной колонкой кракозябр — то есть выгрузка есть, а
    воспользоваться ей нельзя.
    """
    ответ = client.get("/api/analytics/export?period=month&fmt=csv")
    assert ответ.status_code == 200, ответ.text
    assert ответ.headers["content-type"] == "application/zip", ответ.headers

    with zipfile.ZipFile(io.BytesIO(ответ.content)) as архив:
        имена = архив.namelist()
        assert "period.txt" in имена, имена
        таблицы = [и for и in имена if и.endswith(".csv")]
        assert таблицы, имена
        содержимое = архив.read(таблицы[0]).decode("utf-8")
    assert содержимое.startswith("﻿"), "нет метки порядка байтов — Excel даст кракозябры"
    первая = содержимое.splitlines()[0]
    assert ";" in первая, f"разделитель не точка с запятой: {первая!r}"


def test_the_export_names_its_time_zone(client):
    """Отчёт открывают в другом городе.

    Excel про часовые пояса не знает вовсе, поэтому в клетку идёт местное
    время сервера — и молчать об этом нельзя: «14:35» без пояса это число,
    к которому нельзя применить ничего.
    """
    import openpyxl

    ответ = client.get("/api/analytics/export?period=month&fmt=xlsx")
    книга = openpyxl.load_workbook(io.BytesIO(ответ.content))
    ряды = [л for л in книга.worksheets if str(л[1][0].value or "").startswith("Момент")]
    if not ряды:
        pytest.skip("в этой базе нет рядов по времени")
    for лист in ряды:
        подпись = str(лист[1][0].value)
        assert "UTC" in подпись, подпись
        значение = лист.cell(row=2, column=1).value
        assert hasattr(значение, "year"), f"момент записан не датой: {значение!r}"
        assert значение.microsecond == 0, значение


def test_without_openpyxl_the_export_says_so_instead_of_failing(monkeypatch):
    """Пакета может не быть: сервер ставят и в закрытом контуре.

    Пятисотая ошибка вместо отчёта не объясняет ничего, а CSV собирается и
    без единого стороннего пакета.
    """
    from asrhub import analytics_export
    from asrhub.errors import ASRHubError

    настоящий = builtins.__import__

    def без_openpyxl(name, *args, **kwargs):
        if name.startswith("openpyxl"):
            raise ImportError("нет такого пакета")
        return настоящий(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", без_openpyxl)
    with pytest.raises(ASRHubError) as отказ:
        analytics_export.to_xlsx({"overview": {"period": "month"}}, "month")
    assert "openpyxl" in отказ.value.message
    assert "CSV" in (отказ.value.hint or ""), отказ.value.hint

    # А CSV в это же время собирается.
    архив = analytics_export.to_csv_zip({"overview": {"period": "month"}}, "month")
    assert zipfile.ZipFile(io.BytesIO(архив)).namelist()


def test_one_broken_section_does_not_take_the_whole_export_with_it(caplog):
    """У человека должен остаться файл без одного листа, а не ошибка."""
    from asrhub import analytics_export

    отчёт = {
        "overview": {"period": "month", "jobs": {"total": 5}},
        "models": "внезапно строка, а не список",
    }
    with caplog.at_level("WARNING"):
        книга = analytics_export.to_xlsx(отчёт, "month")
    assert книга[:2] == b"PK", "получился не файл xlsx"
    assert any("не попал в выгрузку" in r.message for r in caplog.records), caplog.records


# ---------------------------------------------------------------------------
# Обслуживание по расписанию
# ---------------------------------------------------------------------------

class _Настройки:
    """Настройки для обслуживания без поднятия всего сервера."""

    def __init__(self, каталог: Path, **значения):
        class paths:
            data = str(каталог)
        self.paths = paths
        self._v = значения

    def get(self, key, default=None):
        return self._v.get(key, default)


def _база_с_заданиями(каталог: Path, сколько: int = 20):
    from asrhub.db import Database

    каталог.mkdir(parents=True, exist_ok=True)
    database = Database(каталог / "asrhub.db")
    момент = time.time()
    for i in range(сколько):
        job_id = database.create_job({
            "model": "gigaam-v3-rnnt", "engine": "gigaam", "language": "ru",
            "filename": f"звонок-{i}.wav", "media_duration_s": 300.0,
            "created_at": момент - i * 3600})
        database.update_job(
            job_id, status="failed" if i % 7 == 0 else "completed",
            finished_at=момент - i * 3600 + 60, processing_time_s=44.0,
            queue_time_s=2.0 + i, rtf=0.15, words_count=700, segments_count=40,
            avg_confidence=0.93, device="cuda",
            error_code="decode_error" if i % 7 == 0 else None)
    return database


def test_a_backup_is_a_real_copy_not_a_file_copy(tmp_path: Path):
    """База работает в режиме WAL, рядом лежат -wal и -shm.

    Обычная копия файла получается несогласованной: выглядит как копия и ею
    не является. Глава «Эксплуатация» предупреждает об этом человека —
    сервер обязан делать так же, как советует.
    """
    from asrhub import maintenance
    from asrhub.db import Database

    данные = tmp_path / "data"
    database = _база_с_заданиями(данные, 30)
    настройки = _Настройки(данные, backup_keep=7)

    копия = maintenance.make_backup(database, настройки)
    assert копия is not None and копия.exists(), "копия не сделана"
    assert копия.parent == данные / "backups", копия

    открытая = Database(копия)
    assert открытая.count_jobs() == 30, "в копии не все задания"
    открытая.close()


def test_old_backups_are_removed_but_never_all_of_them(tmp_path: Path):
    """Копия весит примерно столько же, сколько база.

    Без уборки каталог копий растёт быстрее самой базы. Но и обнулять его
    нельзя: последняя копия должна оставаться всегда, каким бы ни было
    значение настройки.
    """
    from asrhub import maintenance

    данные = tmp_path / "data"
    database = _база_с_заданиями(данные, 5)

    for держать, ожидание in ((3, 3), (1, 1), (0, 1)):
        каталог = maintenance.backup_dir(_Настройки(данные))
        for файл in каталог.glob("asrhub-*.db"):
            файл.unlink()
        настройки = _Настройки(данные, backup_keep=держать)
        for i in range(5):
            копия = maintenance.make_backup(database, настройки)
            assert копия is not None
            # Имя копии — по секундам; без сдвига времени пять заходов
            # перезаписали бы один файл, и проверка ничего бы не проверила.
            копия.rename(копия.with_name(f"asrhub-2026010{i}-000000.db"))
            maintenance._подчистить(каталог, держать)
        осталось = list(каталог.glob("asrhub-*.db"))
        assert len(осталось) == ожидание, (держать, [p.name for p in осталось])


def test_restore_checks_the_copy_before_touching_the_live_database(tmp_path: Path):
    """Битый файл, обнаруженный после подмены, оставляет вообще без базы.

    Человек получает два нерабочих файла вместо одного — и это худший
    возможный итог операции, которую затевают ради спасения данных.
    """
    from asrhub import maintenance

    данные = tmp_path / "data"
    данные.mkdir()
    рабочая = данные / "asrhub.db"
    рабочая.write_bytes("рабочая база".encode())
    битая = tmp_path / "битая.db"
    битая.write_bytes("это не база".encode())

    with pytest.raises(ValueError) as отказ:
        maintenance.restore(битая, рабочая)
    assert "не похож на базу" in str(отказ.value), отказ.value
    assert рабочая.read_bytes() == "рабочая база".encode(), "рабочую базу всё же тронули"

    with pytest.raises(FileNotFoundError):
        maintenance.restore(tmp_path / "нет.db", рабочая)


def test_restore_keeps_the_previous_database_and_its_journal(tmp_path: Path):
    """Восстановление не из той копии — обычная ошибка, и она обратима.

    Файлы -wal и -shm обязаны уйти вместе с прежней базой: иначе SQLite
    достроит по ним состояние, которого в восстановленной копии нет.
    """
    from asrhub import maintenance
    from asrhub.db import Database

    данные = tmp_path / "data"
    database = _база_с_заданиями(данные, 12)
    копия = maintenance.make_backup(database, _Настройки(данные))
    assert копия is not None
    database.close()

    рабочая = данные / "asrhub.db"
    рабочая.write_bytes("испорчено".encode())
    Path(str(рабочая) + "-wal").write_bytes("старый журнал".encode())

    maintenance.restore(копия, рабочая)
    assert Database(рабочая).count_jobs() == 12
    assert list(данные.glob("asrhub.db.before-restore-*")), "прежняя база потеряна"
    assert list(данные.glob("asrhub.db-wal.before-restore-*")), "журнал остался на месте"


def test_the_digest_says_the_same_thing_in_words(tmp_path: Path):
    """Приёмник входящих сообщений показывает поле text и ничего больше.

    Без строки словами в чат приходил бы свёрнутый JSON, который никто не
    разворачивает.
    """
    from asrhub import maintenance
    from asrhub.analytics import Analytics

    данные = tmp_path / "data"
    database = _база_с_заданиями(данные, 40)
    сводка = maintenance.build_digest(Analytics(database),
                                      _Настройки(данные, digest_period="month"))
    текст = сводка["text"]
    assert "ASR Hub" in текст
    assert "Заданий: 40" in текст, текст
    assert "—" not in текст.split("Скорость:")[1].split("\n")[0], \
        f"скорость не посчиталась: {текст}"
    assert "уверенность" in текст.lower(), текст
    assert "decode_error" in текст, текст
    assert сводка["instance"], "сводка без имени отправителя"


def test_scheduled_work_waits_its_term_and_does_not_repeat(tmp_path: Path,
                                                           monkeypatch):
    """Отметка «когда в последний раз» лежит в базе и переживает перезапуск.

    Без неё копия делалась бы при каждом старте сервера, а сводка приходила
    бы по разу на перезапуск.
    """
    from asrhub import maintenance
    from asrhub.analytics import Analytics

    данные = tmp_path / "data"
    database = _база_с_заданиями(данные, 5)
    настройки = _Настройки(данные, backup_interval_hours=24, backup_keep=3)
    аналитика = Analytics(database)

    # Первый заход только ставит отметку: сервер только поднялся.
    assert maintenance.run_scheduled(database, настройки, аналитика) == {}
    assert maintenance.run_scheduled(database, настройки, аналитика) == {}

    database.set_kv(maintenance.KV_BACKUP, time.time() - 25 * 3600)
    итог = maintenance.run_scheduled(database, настройки, аналитика)
    assert итог.get("backup"), итог
    # И сразу следом — уже нет.
    assert maintenance.run_scheduled(database, настройки, аналитика) == {}


def test_a_digest_that_cannot_be_delivered_never_breaks_the_server(tmp_path: Path):
    """Сводка — удобство, а не часть обработки заданий."""
    from asrhub import maintenance

    # Адрес, которого нет: порт закрыт.
    assert maintenance.send_digest({"kind": "asrhub.digest"},
                                   "http://127.0.0.1:9/hook") is False
    # И совсем негодный адрес тоже не роняет.
    assert maintenance.send_digest({"kind": "asrhub.digest"}, "не адрес") is False


# ---------------------------------------------------------------------------
# Разграничение доступа и недоверенный ввод
# ---------------------------------------------------------------------------

def test_call_id_cannot_walk_out_of_the_uploads_directory():
    """`call_id` приходит от клиента и становился частью пути.

    Значение вида «../../../имя» уводило запись скачанного файла за пределы
    каталога загрузок, а расширение задавал адрес источника: маршрут
    позволял положить свой .js или .html куда угодно, куда пишет служба, и
    рекурсивно снести чужой каталог при уборке.
    """
    from asrhub.phone_compat import PhoneRequest, safe_path_key

    uploads = Path("/var/lib/asrhub/uploads")
    злые = ["../../../подброшено", "..", "....//..", "a/b/c", "\\\\сервер\\доля",
            ".", "./..", "%2e%2e/", "\x00имя", "id\nимя"]
    for call_id in злые:
        запрос = PhoneRequest(call_id=call_id, part=1,
                              files=["https://x/y.js"], base_url="https://x")
        ключ = запрос.path_key
        assert "/" not in ключ and "\\" not in ключ and ".." not in ключ, (call_id, ключ)
        путь = (uploads / f"{ключ}-часть-0.js").resolve()
        assert uploads.resolve() in путь.parents, (call_id, путь)

    # Обычный идентификатор при этом узнаваем, а не превращён в хеш.
    обычный = PhoneRequest(call_id="CALL-2026-0042", part=2,
                           files=["https://x/y.wav"], base_url="https://x")
    assert обычный.path_key == "CALL-2026-0042_2", обычный.path_key
    # Сам call_id не тронут: он уходит в group_id и в обратный вызов как есть.
    assert обычный.call_id == "CALL-2026-0042"
    assert safe_path_key("") == "без-имени"


def test_deleting_a_job_never_deletes_a_file_outside_the_data_directory(
        client, tmp_path: Path, monkeypatch):
    """Удаление шло по значению из базы без всякой проверки.

    Путь туда кладёт сервер, но стоит ему попасть в базу иначе — правкой
    руками, восстановлением из чужой копии, ошибкой в новом маршруте, — и
    удаление задания превращается в удаление любого файла, до которого
    дотягивается служба.
    """
    посторонний = tmp_path / "чужой-важный-файл.txt"
    посторонний.write_text("не трогать", encoding="utf-8")
    чужой_каталог = tmp_path / "чужой-каталог"
    чужой_каталог.mkdir()
    (чужой_каталог / "внутри.txt").write_text("тоже не трогать", encoding="utf-8")

    state = client.app.state.hub
    job_id = state.db.create_job({
        "filename": "подделка.wav", "model": "demo-simulator",
        "status": "completed",
        "file_path": str(посторонний), "result_path": str(чужой_каталог)})

    ответ = client.delete(f"/api/jobs/{job_id}")
    assert ответ.status_code == 200, ответ.text
    assert посторонний.exists(), "удаление задания снесло посторонний файл"
    assert чужой_каталог.exists(), "удаление задания снесло посторонний каталог"
    assert state.db.get_job(job_id) is None, "само задание должно быть удалено"


def test_cleanup_never_deletes_files_outside_the_data_directory(tmp_path: Path):
    """Та же проверка для уборки по сроку хранения.

    Она идёт раз в час и молча: без проверки это удаление произвольного
    файла по расписанию.
    """
    from asrhub.db import Database

    данные = tmp_path / "data"
    данные.mkdir()
    database = Database(данные / "asrhub.db")

    посторонний = tmp_path / "снаружи.txt"
    посторонний.write_text("не трогать", encoding="utf-8")
    свой = данные / "uploads"
    свой.mkdir()
    внутренний = свой / "своя-запись.wav"
    внутренний.write_bytes(b"x" * 100)

    давно = time.time() - 400 * 86400
    for путь in (посторонний, внутренний):
        job_id = database.create_job({"filename": путь.name, "model": "m",
                                      "file_path": str(путь)})
        database.update_job(job_id, status="completed", finished_at=давно)

    database.cleanup(results_days=30)
    assert посторонний.exists(), "уборка вышла за каталог данных"
    assert not внутренний.exists(), "уборка не удалила свой же файл"


def test_settings_do_not_leak_callback_addresses_and_paths(auth_client, two_users):
    """Входящий адрес чата — это токен, а не просто адрес.

    Кто его знает, тот пишет в чат от имени сервера. Раскладка каталогов —
    разведка перед атакой, и соседний GET /api/system прячет её за правами
    администратора; здесь она уходила любому ключу, что делало ту защиту
    бессмысленной.
    """
    админ = {"X-API-Key": two_users["admin"]}
    обычный = {"X-API-Key": two_users["bob"]}
    секрет = "https://hooks.example.com/services/T0/B0/ОЧЕНЬ-СЕКРЕТНЫЙ-ТОКЕН"

    установка = auth_client.put("/api/settings", json={"webhook_url": секрет},
                                headers=админ)
    assert установка.status_code == 200, установка.text

    чужой = auth_client.get("/api/settings", headers=обычный)
    assert чужой.status_code == 200, чужой.text
    assert секрет not in чужой.text, "адрес обратного вызова ушёл наружу"
    assert чужой.json()["values"]["webhook_url"] == "***"
    assert чужой.json()["paths"] == {}, "раскладка каталогов ушла наружу"

    свой = auth_client.get("/api/settings", headers=админ)
    assert свой.json()["values"]["webhook_url"] == секрет, "администратор своего не видит"
    assert свой.json()["paths"], "администратору раскладка нужна"

    # Ключи доступа и токен не отдаются никому: у токена своя ручка.
    for ответ in (чужой, свой):
        assert "api_keys" not in ответ.json(), ответ.json().keys()
        assert "hf_token" not in ответ.json(), ответ.json().keys()


def test_metric_receivers_hide_their_credentials_from_ordinary_keys(auth_client,
                                                                    two_users):
    """У InfluxDB и Pushgateway пароль стоит прямо в строке запроса.

    Соседние «заменить» и «проверить» требуют администратора, а чтение
    отдавало тот же адрес ключу «только чтение».
    """
    админ = {"X-API-Key": two_users["admin"]}
    только_чтение = {"X-API-Key": two_users["readonly"]}
    адрес = "https://influx.local/write?u=admin&p=ОЧЕНЬ-СЕКРЕТНЫЙ-ПАРОЛЬ"

    установка = auth_client.put("/api/monitoring/targets", headers=админ, json=[
        {"name": "influx", "kind": "influxdb", "url": адрес, "interval_s": 60}])
    if установка.status_code != 200:
        pytest.skip(f"приёмник не завёлся: {установка.text[:120]}")

    чужой = auth_client.get("/api/monitoring/targets", headers=только_чтение)
    assert чужой.status_code == 200, чужой.text
    assert "СЕКРЕТНЫЙ" not in чужой.text, "учётные данные приёмника ушли наружу"
    # Но куда шлём — видно: «***» на этот вопрос не отвечает.
    assert "influx.local" in чужой.text, чужой.text

    свой = auth_client.get("/api/monitoring/targets", headers=админ)
    assert адрес in свой.text, "администратор своего адреса не видит"


def test_analytics_works_for_a_key_that_belongs_to_a_group(tmp_path: Path):
    """Ключ в подразделении видит задания всей группы.

    `scope_owner` отдаёт для него перечень владельцев, а список нехешируем:
    общий кеш выборок падал с TypeError, и такой ключ получал 500 на всей
    аналитике — ровно там, где разграничение и работает.
    """
    from asrhub.analytics import Analytics
    from asrhub.db import Database

    database = Database(tmp_path / "g.sqlite3")
    момент = time.time()
    for кто in ("alice", "bob", "carol"):
        job_id = database.create_job({"model": "m", "owner": кто,
                                      "media_duration_s": 60.0,
                                      "created_at": момент})
        database.update_job(job_id, status="completed", finished_at=момент + 10,
                            processing_time_s=10.0)

    отчёт = Analytics(database).full_report("month", owner=["alice", "bob"])
    assert отчёт["overview"]["jobs"]["total"] == 2, отчёт["overview"]["jobs"]
    # И одиночный владелец не сломался заодно.
    один = Analytics(database).full_report("month", owner="carol")
    assert один["overview"]["jobs"]["total"] == 1, один["overview"]["jobs"]


def test_the_snippet_search_is_limited_to_the_rows_being_shown(tmp_path: Path):
    """Предел выборки съедали чужие совпадения.

    Поиск шёл по всей таблице реплик, и только потом результат отсеивался по
    показанным заданиям: на оживлённом архиве фраза пропадала из строки тем
    чаще, чем активнее соседи, — и без всякой видимой причины.
    """
    from asrhub.db import Database

    database = Database(tmp_path / "n.sqlite3")
    for i in range(400):
        job_id = database.create_job({"model": "m", "owner": "сосед",
                                      "filename": f"ч-{i}.wav"})
        database.save_segments(job_id, [
            {"start": 0.0, "end": 5.0, "text": "обсудили договор поставки"}])
    свои = []
    for i in range(3):
        job_id = database.create_job({"model": "m", "owner": "мы",
                                      "filename": f"с-{i}.wav"})
        database.save_segments(job_id, [
            {"start": 7.0, "end": 12.0, "text": "тоже про договор и сроки"}])
        свои.append(job_id)

    находки = database.best_snippets("договор", свои)
    assert set(находки) == set(свои), (
        f"фраза нашлась только для {len(находки)} из {len(свои)} показанных строк")
    for находка in находки.values():
        assert находка["start_s"] == 7.0, находка


# ---------------------------------------------------------------------------
# Надёжность: чужие задания, ресурсы, служебный цикл
# ---------------------------------------------------------------------------

def test_a_failure_never_overwrites_a_job_taken_over_by_a_neighbour(tmp_path: Path):
    """Успешное завершение было защищено, а три ветки отказа — нет.

    Экземпляр, застрявший дольше отметки жизни, оживал и писал свой отказ
    поверх задания, которое уже считает сосед. Когда сосед досчитывал, его
    защищённая запись не проходила, и **готовая расшифровка выбрасывалась**:
    пользователь получал «ошибка» вместо результата.
    """
    import threading

    from asrhub import job_queue as JQ
    from asrhub.db import Database

    database = Database(tmp_path / "q.sqlite3")
    очередь = JQ.JobQueue.__new__(JQ.JobQueue)
    очередь.db = database
    очередь._lock = threading.RLock()
    очередь._cancelled = set()
    очередь._discard_results = lambda *a: None

    job_id = database.create_job({"model": "m", "filename": "разговор.wav"})
    database.update_job(job_id, status=JQ.STATUS_RUNNING,
                        instance_id=JQ.INSTANCE_ID, heartbeat_at=1.0)
    assert очередь._write_own(job_id, status=JQ.STATUS_FAILED, finished_at=3.0), \
        "своё задание записать не дали"

    # Сосед забрал зависшее задание себе и считает.
    database.update_job(job_id, status=JQ.STATUS_RUNNING, instance_id="сосед",
                        heartbeat_at=9.0, retries=1)
    assert not очередь._write_own(job_id, status=JQ.STATUS_FAILED, finished_at=4.0), \
        "отказ записался поверх задания соседа"

    итог = database.get_job(job_id)
    assert итог["status"] == JQ.STATUS_RUNNING, итог["status"]
    assert итог["instance_id"] == "сосед", итог["instance_id"]
    assert итог["retries"] == 1, "счётчик повторов затёрт устаревшим снимком"


def test_finished_jobs_do_not_pile_up_in_the_concurrency_table():
    """Освобождение слота было написано дважды, и копии разошлись.

    Вторая снимала слот и счётчик модели, но не отметку одновременности —
    та копилась по записи на каждое проведённое задание. Мало того что без
    предела: этот словарь обходится под общей блокировкой при каждом старте
    задания, и на сотне тысяч выходило девять миллисекунд блокировки на
    задание. Со стороны — «сервер к вечеру тупеет».
    """
    import threading

    from asrhub.job_queue import JobQueue

    очередь = JobQueue.__new__(JobQueue)
    очередь._lock = threading.RLock()
    очередь._running, очередь._concurrency, очередь._model_counts = {}, {}, {}

    for i in range(50):
        job_id = f"job{i}"
        очередь._running[job_id] = 0.0
        очередь._concurrency[job_id] = 0
        очередь._model_counts["m"] = очередь._model_counts.get("m", 0) + 1
        очередь._note_concurrency()
        очередь._release_slot(job_id, "m")

    assert очередь._running == {}, очередь._running
    assert очередь._concurrency == {}, f"накопилось {len(очередь._concurrency)} записей"
    assert очередь._model_counts == {"m": 0}, очередь._model_counts

    # И освобождение живёт в одном месте, а не в двух.
    источник = Path(__file__).resolve().parent.parent / "server" / "asrhub" / "job_queue.py"
    текст = источник.read_text(encoding="utf-8")
    начало = текст.index("            finally:\n                # Освобождение")
    assert "_release_slot(" in текст[начало:начало + 800], "копия освобождения вернулась"


def test_backups_keep_working_on_a_disk_that_is_almost_full(tmp_path: Path):
    """Подчистка стояла после копирования, и до неё не доходило дело.

    Каталог копий по умолчанию лежит на той же файловой системе, что и база.
    Когда места хватало ровно на `backup_keep` копий, очередная не
    помещалась, старые не удалялись — и свежих копий не появлялось больше
    никогда, молча.
    """
    from asrhub import maintenance

    данные = tmp_path / "data"
    database = _база_с_заданиями(данные, 5)
    каталог = maintenance.backup_dir(_Настройки(данные))
    настройки = _Настройки(данные, backup_keep=2)

    имена = []
    for i in range(5):
        копия = maintenance.make_backup(database, настройки)
        assert копия is not None, f"копия {i + 1} не сделана"
        имена.append(копия.name)
        # Имя с точностью до секунды: без сдвига пять заходов легли бы в
        # один файл, и проверка ничего бы не проверила.
        копия.rename(копия.with_name(f"asrhub-2026010{i}-000000-x.db"))

    осталось = sorted(p.name for p in каталог.glob("asrhub-*.db"))
    assert len(осталось) == 2, осталось
    # Именно последние, а не первые попавшиеся.
    assert осталось == ["asrhub-20260103-000000-x.db",
                        "asrhub-20260104-000000-x.db"], осталось
    # Имя несёт экземпляр: на общей базе два сервера в одну секунду выбирали
    # одно имя, и вместо двух копий оставалась одна.
    assert all(len(и.split("-")) >= 4 for и in имена), имена


def test_lowering_and_raising_the_worker_count_gets_the_workers_back(client):
    """Помеченный воркер уходит не сразу — он замечает пометку, проснувшись.

    Уменьшить и тут же вернуть обратно означало, что ни одна ветка не
    сработала: число воркеров ещё прежнее, и пометка оставалась. Сервер жил
    с половиной заявленных до перезапуска, показывая в состоянии полное
    число.
    """
    import threading

    очередь = client.app.state.hub.queue

    def живых() -> int:
        return sum(1 for t in threading.enumerate()
                   if t.name.startswith("asrhub-worker-") and t.is_alive())

    было = живых()
    assert было >= 2, f"для проверки нужно хотя бы два воркера, а их {было}"

    очередь.set_concurrency(1)
    time.sleep(1.2)
    очередь.set_concurrency(было)
    time.sleep(2.5)

    assert not очередь._retiring, f"пометка на выход осталась: {sorted(очередь._retiring)}"
    assert живых() == было, f"вернулось {живых()} воркеров из {было}"


def test_a_cached_result_carries_every_counter_the_analytics_reads(tmp_path: Path):
    """Строка «Знаков» занижалась ровно на долю попаданий в кеш.

    При том что «Слов» считалось полностью, так что расхождение выглядело
    как ошибка в подсчёте, а не как пропуск. Ноль говорящих вдобавок молча
    выбрасывал такие задания из разреза по числу собеседников.
    """
    источник = (Path(__file__).resolve().parent.parent / "server" / "asrhub"
                / "job_queue.py").read_text(encoding="utf-8")
    начало = источник.index("    def _clone_cached(")
    тело = источник[начало:начало + 3200]
    for поле in ("words_count", "chars_count", "segments_count",
                 "speakers_count", "avg_confidence", "rtf"):
        assert f"{поле}=cached.get(" in тело, f"из кеша не переносится {поле}"


def test_a_broken_janitor_step_does_not_flood_the_log(tmp_path: Path, caplog):
    """Ограничитель записей не включался никогда.

    Сбой шага ловился вложенным `except`, внешний try завершался штатно, и
    ветка «иначе» объявляла восстановление и обнуляла счётчик. В журнал
    каждые двадцать секунд шла пара строк «дал сбой (1-й раз)» и
    «восстановился» — ровно то заливание, против которого ограничитель и
    написан.
    """
    import threading

    from asrhub import job_queue as JQ
    from asrhub.db import Database

    очередь = JQ.JobQueue.__new__(JQ.JobQueue)
    очередь.db = Database(tmp_path / "j.sqlite3")
    очередь.settings = _Настройки(tmp_path)
    очередь._stop = threading.Event()
    очередь._lock = threading.RLock()
    очередь._running, очередь._concurrency = {}, {}
    очередь._sample_system = lambda: None
    очередь._reclaim_stale_jobs = lambda: None
    очередь._analytics = lambda: None

    class Реестр:
        @staticmethod
        def collect_idle() -> None:
            raise RuntimeError("датчик недоступен")

    очередь.registry = Реестр()

    прежний = JQ.SAMPLE_PERIOD_S
    JQ.SAMPLE_PERIOD_S = 0.03
    try:
        with caplog.at_level("INFO", logger="asrhub.queue"):
            поток = threading.Thread(target=очередь._janitor_loop, daemon=True)
            поток.start()
            time.sleep(0.8)
            очередь._stop.set()
            поток.join(timeout=3)
    finally:
        JQ.SAMPLE_PERIOD_S = прежний

    записи = [r.getMessage() for r in caplog.records]
    сбои = [r for r in записи if "дал сбой" in r]
    assert 0 < len(сбои) <= 4, f"ограничитель не работает: {len(сбои)} записей о сбое"
    assert not [r for r in записи if "восстановился" in r], \
        "цикл объявил восстановление, хотя шаг падает по-прежнему"


def test_the_tag_breakdown_has_a_ceiling(tmp_path: Path):
    """Метку задаёт клиент, и её мощность ничем не ограничена.

    Клиент, ставящий уникальную метку на каждое задание, превращал разрез в
    сто тысяч строк — и в таблице на экране, и в листе выгрузки, где
    остальные разрезы не длиннее полусотни.
    """
    from asrhub.analytics import TAG_LIMIT, Analytics
    from asrhub.db import Database

    database = Database(tmp_path / "t.sqlite3")
    момент = time.time()
    for i in range(TAG_LIMIT * 3):
        job_id = database.create_job({"model": "m", "media_duration_s": 60.0,
                                      "tags": f"метка-{i}", "created_at": момент})
        database.update_job(job_id, status="completed", finished_at=момент + 5,
                            processing_time_s=5.0)

    метки = Analytics(database).by_tag("month")
    assert len(метки) == TAG_LIMIT, len(метки)
    # Верх списка — по числу заданий, а не по алфавиту.
    assert метки == sorted(метки, key=lambda r: -r["jobs"]), метки[:3]
